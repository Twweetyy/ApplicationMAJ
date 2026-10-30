"""Onglet « Portrait Roubaix » : chiffres clés de la ville, juste après la Présentation.

Chaque valeur est une formule Excel qui pointe vers les onglets thématiques du classeur :
- effectif : somme des IRIS ;
- part / taux : somme du numérateur / somme du dénominateur (jamais une moyenne de taux) ;
- Filosofi (médianes, taux de pauvreté) : non additif, on donne l'IRIS le plus bas et le plus haut.
Un indicateur dont l'onglet ou la colonne manque est simplement omis.
"""
import re
import unicodedata

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NOM_ONGLET = "Portrait Roubaix"
BLEU, GRIS, BLANC = "1F4E78", "595959", "FFFFFF"
FORMATS = {"entier": "#,##0", "pct": "0.0 %", "euro": '#,##0" €"', "dec2": "0.00", "pct100": '0.0" %"'}
MOTIF_ANNEE = re.compile(r"^(.*\S)\s+(\d{4})$")

# Source affichée sous chaque indicateur, par thème d'onglet
SOURCES = {
    "POP": "RP", "ACT RES": "RP", "COU FAM": "RP", "DIP FORMATION": "RP", "LOG": "RP",
    "PETITE ENFANCE": "RP (fichier détail)", "REV DEC": "Filosofi", "REV DISP": "Filosofi",
    "CAF": "CAF", "CNAM": "CNAM", "DEFM": "France Travail", "RPLS": "RPLS",
}

# (titre, couleur, couleur claire, description)
BLOCS = [
    ("Démographie", "2E75B6", "DEEAF6",
     "Population, densité, âges, sexe, nationalité, ménages et familles (monoparentales, nombreuses)."),
    ("Emploi et formation", "C55A11", "FBE5D6",
     "Activité, emploi, chômage, diplômes de la population non scolarisée et demandeurs d'emploi."),
    ("Revenus et prestations", "7030A0", "EADCF4",
     "Revenus et pauvreté (Filosofi), allocataires CAF, RSA, AAH et complémentaire santé solidaire (CNAM)."),
    ("Logement", "548235", "E2F0D9",
     "Parc, vacance, typologie, statut d'occupation, logements sociaux, ancienneté et suroccupation."),
]

# Indicateurs : (bloc, libellé, onglet, calcul, colonnes, format)
#   somme : colonnes = [colonne]           ratio : colonnes = ([numérateurs], [dénominateurs])
#   min / max : colonnes = [colonne]       (valeur de l'IRIS le plus bas / le plus haut)
INDICATEURS = [
    ("Démographie", "Habitants", "POP", "somme", ["Population"], "entier"),
    ("Démographie", "Habitants au km²", "POP", "ratio", (["Population"], ["Superficie en km²"]), "entier"),
    ("Démographie", "Moins de 20 ans", "POP", "ratio", (["Pop 0-19 ans"], ["Population"]), "pct"),
    ("Démographie", "65 ans ou plus", "POP", "ratio", (["Pop 65 ans ou plus"], ["Population"]), "pct"),
    ("Démographie", "Femmes", "POP", "ratio", (["Pop Femmes"], ["Population"]), "pct"),
    ("Démographie", "Étrangers", "POP", "ratio", (["Pop Etrangers"], ["Population"]), "pct"),
    ("Démographie", "Immigrés", "POP", "ratio", (["Pop Immigrés"], ["Population"]), "pct"),
    ("Démographie", "Enfants de 0 à 17 ans", "PETITE ENFANCE", "somme", ["TOTAL 0-17"], "entier"),
    ("Démographie", "Ménages", "COU FAM", "somme", ["Ménages"], "entier"),
    ("Démographie", "Taille moyenne des ménages", "COU FAM", "ratio", (["Pop Ménages"], ["Ménages"]), "dec2"),
    ("Démographie", "Familles monoparentales", "COU FAM", "ratio", (["Fam Monoparentales"], ["Familles"]), "pct"),
    ("Démographie", "Familles nombreuses (3 enfants ou plus)", "COU FAM", "ratio",
     (["Fam 3 enfants moins 25 ans", "Fam 4 enfants ou plus moins 25 ans"], ["Familles"]), "pct"),

    ("Emploi et formation", "Actifs (15-64 ans)", "ACT RES", "somme", ["Actifs 15-64 ans"], "entier"),
    ("Emploi et formation", "Taux d'emploi (15-64 ans)", "ACT RES", "ratio",
     (["Actifs occupés 15-64 ans"], ["Pop 15-64 ans"]), "pct"),
    ("Emploi et formation", "Taux de chômage (15-64 ans)", "ACT RES", "ratio",
     (["Chômeurs 15-64 ans"], ["Actifs 15-64 ans"]), "pct"),
    ("Emploi et formation", "Taux de chômage des 15-24 ans", "ACT RES", "ratio",
     (["Chômeurs 15-24 ans"], ["Actifs 15-24 ans"]), "pct"),
    ("Emploi et formation", "Salariés en CDI ou fonction publique", "ACT RES", "ratio",
     (["Salariés 15 ans ou plus Fonct publ, CDI"], ["Actifs occupés 15 ans ou plus"]), "pct"),
    ("Emploi et formation", "Sans diplôme ou CEP (15 ans ou plus non scolarisés)", "DIP FORMATION", "ratio",
     (["Pop 15 ans ou plus non scol. Sans diplôme ou CEP"], ["Pop 15 ans ou plus non scolarisée"]), "pct"),
    ("Emploi et formation", "Diplômés du supérieur (15 ans ou plus non scolarisés)", "DIP FORMATION", "ratio",
     ([
         "Pop 15 ans ou plus non scol. Enseignement sup de niveau bac + 2",
         "Pop 15 ans ou plus non scol. Enseignement sup de niveau bac + 3 ou 4",
         "Pop 15 ans ou plus non scol. Enseignement sup de niveau bac + 5 ou plus",
     ], ["Pop 15 ans ou plus non scolarisée"]), "pct"),
    ("Emploi et formation", "Demandeurs d'emploi (cat. ABC)", "DEFM", "somme", ["DEFM catégories ABC"], "entier"),
    ("Emploi et formation", "Demandeurs ABC de moins de 26 ans", "DEFM", "ratio",
     (["DEFM catégories ABC moins de 26 ans"], ["DEFM catégories ABC"]), "pct"),
    ("Emploi et formation", "Demandeurs ABC depuis 2 ans ou plus", "DEFM", "ratio",
     (["DEFM catégories ABC d'ancienneté au chômage d'au moins 2 ans"], ["DEFM catégories ABC"]), "pct"),

    ("Revenus et prestations", "Revenu disponible médian : IRIS le plus bas", "REV DISP", "min", ["Médiane (€)"], "euro"),
    ("Revenus et prestations", "Revenu disponible médian : IRIS le plus haut", "REV DISP", "max", ["Médiane (€)"], "euro"),
    ("Revenus et prestations", "Taux de pauvreté : IRIS le plus bas", "REV DISP", "min",
     ["Taux de pauvreté au seuil de 60 % (%)"], "pct100"),
    ("Revenus et prestations", "Taux de pauvreté : IRIS le plus haut", "REV DISP", "max",
     ["Taux de pauvreté au seuil de 60 % (%)"], "pct100"),
    ("Revenus et prestations", "Allocataires CAF", "CAF", "somme", ["Nombre total d'allocataires"], "entier"),
    ("Revenus et prestations", "Population couverte par la CAF", "CAF", "ratio",
     (["Personnes couvertes"], ["Population (CAF)"]), "pct"),
    ("Revenus et prestations", "Allocataires CAF au RSA socle", "CAF", "ratio",
     (["Allocataires percevant le RSA socle"], ["Nombre total d'allocataires"]), "pct"),
    ("Revenus et prestations", "Allocataires CAF mono-parents", "CAF", "ratio",
     (["Allocataires mono-parent"], ["Nombre total d'allocataires"]), "pct"),
    ("Revenus et prestations", "Allocataires de l'AAH", "CAF", "somme",
     ["Allocataires percevant l’Allocation Adulte Handicapé"], "entier"),
    ("Revenus et prestations", "Bénéficiaires de la C2S non participative", "CNAM", "ratio",
     (["Population couverte bénéficiaire de la C2S non participative"], ["Population couverte"]), "pct"),
    ("Revenus et prestations", "Bénéficiaires de la C2S participative", "CNAM", "ratio",
     (["Population couverte bénéficiaire de la C2S participative"], ["Population couverte"]), "pct"),

    ("Logement", "Logements", "LOG", "somme", ["Logements"], "entier"),
    ("Logement", "Résidences principales", "LOG", "somme", ["Résidences principales"], "entier"),
    ("Logement", "Logements vacants", "LOG", "ratio", (["Logements vacants"], ["Logements"]), "pct"),
    ("Logement", "Maisons", "LOG", "ratio", (["Maisons"], ["Logements"]), "pct"),
    ("Logement", "Propriétaires occupants", "LOG", "ratio",
     (["Rés princ occupées Propriétaires"], ["Résidences principales"]), "pct"),
    ("Logement", "Locataires HLM", "LOG", "ratio", (["Rés princ HLM louée vide"], ["Résidences principales"]), "pct"),
    ("Logement", "Résidences construites avant 1946", "LOG", "ratio",
     (["Rés princ avt 1919", "Rés princ 1919 à 1945"], ["Rés princ avt 2020"]), "pct"),
    ("Logement", "Résidences en suroccupation", "LOG", "ratio",
     ([
         "résidences principales en suroccupation modérée",
         "résidences principales en suroccupation accentuée",
     ], ["Résidences principales"]), "pct"),
    ("Logement", "Ménages avec au moins une voiture", "LOG", "ratio",
     (["Ménages au moins une voiture"], ["Ménages"]), "pct"),
    ("Logement", "Logements sociaux (RPLS)", "RPLS", "somme", ["Nombre de logements du Parc Locatif Social"], "entier"),
]

NOTES = [
    "Les parts et taux sont recalculés sur Roubaix entière (somme des IRIS), pas en moyenne des IRIS.",
    "Les revenus Filosofi (médianes, taux de pauvreté) ne s'additionnent pas : l'IRIS le plus bas et le plus haut sont donnés.",
    "Les valeurs sont des formules : elles suivent les onglets du classeur. Cliquer sur une source ouvre l'onglet.",
]


def _normaliser(texte) -> str:
    texte = unicodedata.normalize("NFKD", str(texte).casefold())
    texte = "".join(c for c in texte if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]", "", texte)


def _onglets_par_theme(workbook) -> dict:
    """{'POP': 'POP 2023', 'PETITE ENFANCE': 'PETITE ENFANCE 2022', ...}"""
    resultat = {}
    for nom in workbook.sheetnames:
        correspondance = MOTIF_ANNEE.match(nom)
        if correspondance:
            resultat.setdefault(correspondance.group(1), nom)
    return resultat


def _plages(ws) -> dict:
    """{en-tête normalisé: "'POP 2023'!$C$3:$C$40"} pour les lignes d'IRIS de l'onglet."""
    derniere = ws.max_row
    while derniere > 3 and all(c.value in (None, "") for c in ws[derniere]):
        derniere -= 1
    plages = {}
    for cellule in ws[2]:
        if cellule.value not in (None, ""):
            lettre = get_column_letter(cellule.column)
            plages.setdefault(_normaliser(cellule.value), f"'{ws.title}'!${lettre}$3:${lettre}${derniere}")
    return plages


def _formule(calcul: str, colonnes, plages: dict) -> str | None:
    def sommes(noms):
        refs = [plages.get(_normaliser(nom)) for nom in noms]
        return None if None in refs else "+".join(f"SUM({ref})" for ref in refs)

    if calcul == "ratio":
        numerateur, denominateur = sommes(colonnes[0]), sommes(colonnes[1])
        if numerateur is None or denominateur is None:
            return None
        return f'=IFERROR(({numerateur})/({denominateur}),"")'
    ref = plages.get(_normaliser(colonnes[0]))
    if ref is None:
        return None
    fonction = {"somme": "SUM", "min": "MIN", "max": "MAX"}[calcul]
    return f"={fonction}({ref})"


def _fusion(ws, plage, valeur=None, font=None, fill=None, align=None):
    ws.merge_cells(plage)
    cellule = ws[plage.split(":")[0]]
    cellule.value = valeur
    if fill:
        for ligne in ws[plage]:
            for c in ligne:
                c.fill = fill
    if font:
        cellule.font = font
    cellule.alignment = align or Alignment(vertical="center", wrap_text=True)
    return cellule


def ajouter_portrait_roubaix(workbook) -> None:
    """Crée (ou recrée) l'onglet « Portrait Roubaix » juste après « Présentation »."""
    if NOM_ONGLET in workbook.sheetnames:
        del workbook[NOM_ONGLET]
    onglets = _onglets_par_theme(workbook)
    if not onglets:
        return
    position = workbook.sheetnames.index("Présentation") + 1 if "Présentation" in workbook.sheetnames else 0
    ws = workbook.create_sheet(NOM_ONGLET, position)
    ws.sheet_view.showGridLines = False
    for i, largeur in enumerate([2] + [32, 12, 3] * len(BLOCS), start=1):
        ws.column_dimensions[get_column_letter(i)].width = largeur
    fin = get_column_letter(1 + 3 * len(BLOCS) - 1)

    annee = MOTIF_ANNEE.match(onglets.get("POP", "")) if "POP" in onglets else None
    titre = f"Portrait de Roubaix · {annee.group(2)}" if annee else "Portrait de Roubaix"
    _fusion(ws, f"B2:{fin}2", titre, Font(bold=True, size=20, color=BLEU))
    ws.row_dimensions[2].height = 34
    _fusion(ws, f"B3:{fin}3",
            "Chiffres clés de Roubaix (ensemble des IRIS), toutes sources du classeur confondues. "
            "Chaque indicateur indique sa source et son année.",
            Font(italic=True, size=10, color=GRIS))
    ws.row_dimensions[3].height = 22

    plages = {theme: _plages(workbook[nom]) for theme, nom in onglets.items()}
    filet = Border(top=Side(style="hair", color="BFBFBF"))
    hauteur_max = 8
    for k, (bloc, couleur, clair, description) in enumerate(BLOCS):
        c1, c2 = get_column_letter(2 + 3 * k), get_column_letter(3 + 3 * k)
        _fusion(ws, f"{c1}5:{c2}6", bloc.upper(), Font(bold=True, size=14, color=BLANC),
                PatternFill("solid", fgColor=couleur), Alignment(horizontal="center", vertical="center"))
        _fusion(ws, f"{c1}7:{c2}7", description, Font(size=9, color=GRIS),
                PatternFill("solid", fgColor=clair), Alignment(wrap_text=True, vertical="top"))
        ligne = 8
        for bloc_indicateur, libelle, theme, calcul, colonnes, fmt in INDICATEURS:
            if bloc_indicateur != bloc or theme not in onglets:
                continue
            formule = _formule(calcul, colonnes, plages[theme])
            if formule is None:
                continue
            a, b = ws[f"{c1}{ligne}"], ws[f"{c2}{ligne}"]
            a.value, a.font = libelle, Font(size=10)
            a.alignment = Alignment(wrap_text=True, vertical="center")
            b.value, b.number_format = formule, FORMATS[fmt]
            b.font = Font(bold=True, size=13, color=couleur)
            b.alignment = Alignment(horizontal="right", vertical="center")
            a.border = b.border = filet
            nom_onglet = onglets[theme]
            source = ws[f"{c1}{ligne + 1}"]
            source.value = f"{SOURCES.get(theme, theme)} {MOTIF_ANNEE.match(nom_onglet).group(2)} · onglet {nom_onglet}"
            source.hyperlink = f"#'{nom_onglet}'!A1"
            source.font = Font(size=8, color=GRIS, underline="single")
            ws.row_dimensions[ligne].height = 30
            ligne += 2
        hauteur_max = max(hauteur_max, ligne)
    ws.row_dimensions[7].height = 46

    r = hauteur_max + 1
    _fusion(ws, f"B{r}:{fin}{r}", "À savoir", Font(bold=True, color=BLANC), PatternFill("solid", fgColor=GRIS))
    for i, note in enumerate(NOTES, start=1):
        _fusion(ws, f"B{r + i}:{fin}{r + i}", "• " + note, Font(size=9))
        ws.row_dimensions[r + i].height = 18
    ws.sheet_view.topLeftCell = "A1"
