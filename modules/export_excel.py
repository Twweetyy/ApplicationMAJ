import os
import re
import unicodedata
from copy import copy
import pandas as pd
import openpyxl
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.table import TableColumn
import yaml
from tqdm.auto import tqdm
from .pages_info import mettre_a_jour_pages_info
from .mise_en_page import appliquer_mise_en_page
from .portrait_roubaix import ajouter_portrait_roubaix

# En-têtes du modèle DEFM renommés selon les libellés des données (clé = en-tête normalisé)
ALIAS_ENTETES_DEFM = {
    "nombre total d allocataires": "DEFM catégorie A",
    "defm categories abc dont motif d inscription est une 1ere entree sur marche du travail":
        "DEFM catégories ABC dont le motif d'inscription est une 1ère entrée sur le marché du travail",
    "defm categories abc dont motif d inscription est un licenciement":
        "DEFM catégories ABC dont le motif d'inscription est un licenciement",
    "defm categories abc dont motif d inscription est une fin de contrat a duree determinee":
        "DEFM catégories ABC dont le motif d'inscription est une fin de CDD",
    "defm categories abc dont motif d inscription est une fin de mission d interim":
        "DEFM catégories ABC dont le motif d'inscription est une fin de mission d'intérim",
    "defm categories abc dont motif d inscription est autre":
        "DEFM catégories ABC dont le motif d'inscription est autre",
}

# Seules colonnes de zonage (grises) conservées dans les onglets du classeur, dans cet ordre.
# « Grands quartiers » vient de grand_quartier.yml (via Périmètres_MAS.csv), le libellé IRIS reste celui de l'INSEE.
COLONNES_ZONAGE = ["Mairies de quartiers", "Grands quartiers", "Libellé de l'Iris"]
# Libellés de ces colonnes dans le modèle petite enfance
ZONAGE_PETITE_ENFANCE = {
    "Mairies de quartiers": "Mairies de qartiers",
    "Grands quartiers": "Quartiers SIG (Grands quartiers)",
}
# Colonnes de Périmètres_MAS.csv retirées du classeur (gardées pour les exports SIG / SQL)
COLONNES_ZONAGE_RETIREES = [
    "Centres sociaux",
    "Périmètres NPNRU",
    "Contrats de Ville 2024 (CDV)",
    "Cités éducatives",
]

# Libellés CNAM des données -> en-têtes du modèle
RENOMMAGES_CNAM = {
    "Population couverte bénéficiaire du régime général": "Population couverte",
    "Hommes bénéficiaires du régime général": "Hommes",
    "Femmes bénéficiaires du régime général": "Femmes",
    "Bénéficiaires du régime général de moins de 18 ans": "Moins de 18 ans",
    "Bénéficiaires du régime général  de 18 ans à moins de 25 ans": "de 18 ans à moins de 25 ans",
    "Bénéficiaires du régime général  de 25 ans à moins de 50 ans": "de 25 ans à moins de 50 ans",
    "Bénéficiaires du régime général  de 50 ans à moins de 60 ans": "de 50 ans à moins de 60 ans",
    "Bénéficiaires du régime général  de 60 ans à moins de 75 ans": "de 60 ans à moins de 75 ans",
    "Bénéficiaires du régime général  de 75 ans ou plus": "75 ans ou plus",
}


def charger_titres_onglets(chemin_yaml: str) -> dict:
    """Charge le dictionnaire contenant les grands titres de chaque onglet."""
    if not os.path.exists(chemin_yaml):
        return {}
    with open(chemin_yaml, "r", encoding="UTF-8") as file:
        return yaml.safe_load(file) or {}

def titre_onglet(titres: dict, nom_onglet: str) -> str:
    """Titre d'un onglet quel que soit le millésime (clé 'POP' ou 'POP 2021' -> onglet 'POP 2023')."""
    if nom_onglet in titres:
        return titres[nom_onglet]
    theme = re.sub(r"\s+\d{4}$", "", nom_onglet)
    for cle, titre in titres.items():
        if re.sub(r"\s+\d{4}$", "", str(cle)) == theme:
            return titre
    return f"Données infra-communales - {nom_onglet}"

def formater_onglet_thematique(worksheet, dataframe, titre_onglet):
    """
    Applique la charte graphique Mairie/MEL sur un onglet de données :
    - Ligne 1 : Titre principal (A1) et Sous-totaux dynamiques (Vert).
    - Ligne 2 : En-têtes (Gris pour Géo, Bleu pour Insee, Orange pour Calculs).
    """
    # --- DÉFINITION DES STYLES ---
    font_bold = Font(bold=True)
    font_total = Font(bold=True, color="00B050")
    font_titre = Font(bold=True, color="800000", size=12)
    
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center")
    
    fill_geo = PatternFill(start_color="BFBFBF", end_color="BFBFBF", fill_type="solid")
    fill_data_bleu = PatternFill(start_color="91DEFF", end_color="91DEFF", fill_type="solid")
    fill_data_orange = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
    fill_total_vert = PatternFill(start_color="EBF1DE", end_color="EBF1DE", fill_type="solid")
    fill_gris_clair = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

    border_gris = Border(
        left=Side(style="thin", color="BFBFBF"), right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"), bottom=Side(style="thin", color="BFBFBF")
    )

    # Paramétrage des dimensions
    worksheet.row_dimensions[1].height = 30
    worksheet.row_dimensions[2].height = 80

    # Détection des séparations de colonnes (Géo vs Données vs Calculs)
    colonnes = list(dataframe.columns)
    derniere_geo = "IRIS" if "IRIS" in colonnes else "Libellé de l'Iris"
    col_geo_max = colonnes.index(derniere_geo) + 1 if derniere_geo in colonnes else len(COLONNES_ZONAGE)
    
    # On cherche "Superficie en km²" ou la fin des données brutes pour appliquer l'orange
    try:
        col_calc_start = dataframe.columns.get_loc("Superficie en km²") + 1
    except KeyError:
        col_calc_start = worksheet.max_column - 5 # Approximation si la colonne n'existe pas

    # --- ÉCRITURE DE LA LIGNE 1 (Titre et Sous-Totaux) ---
    worksheet.cell(row=1, column=1, value=titre_onglet).font = font_titre
    worksheet.cell(row=1, column=1).alignment = align_left
    
    worksheet.cell(row=1, column=col_geo_max, value="Total").font = font_bold
    worksheet.cell(row=1, column=col_geo_max).alignment = align_center

    for col_idx in range(col_geo_max + 1, worksheet.max_column + 1):
        lettre_col = get_column_letter(col_idx)
        # Formule dynamique prenant en compte les filtres utilisateurs
        formule = f"=SUBTOTAL(9,{lettre_col}3:{lettre_col}{worksheet.max_row})"
        
        cell_total = worksheet.cell(row=1, column=col_idx, value=formule)
        cell_total.font = font_total
        cell_total.fill = fill_total_vert
        cell_total.alignment = align_center
        cell_total.border = border_gris

    # --- APPLICATION DES STYLES POUR LE RESTE DU TABLEAU ---
    for row_idx in range(2, worksheet.max_row + 1):
        for col_idx in range(1, worksheet.max_column + 1):
            cell = worksheet.cell(row=row_idx, column=col_idx)
            cell.border = border_gris
            
            # Ligne 2 (En-têtes)
            if row_idx == 2:
                cell.alignment = align_center
                cell.font = font_bold
                
                # Ajustement largeur (spécifique pour le libellé IRIS)
                largeur = 24 if cell.value == "Libellé de l'Iris" else 15
                worksheet.column_dimensions[get_column_letter(col_idx)].width = largeur
                
                # Couleurs des en-têtes
                if col_idx <= col_geo_max:
                    cell.fill = fill_geo
                elif col_idx < col_calc_start:
                    cell.fill = fill_data_bleu
                else:
                    cell.fill = fill_data_orange
            
            # Lignes de données (3 et +)
            else:
                if col_idx <= col_geo_max:
                    cell.font = font_bold
                    cell.fill = fill_gris_clair
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="right", vertical="center")

    # Figeage des volets (sous les en-têtes, après le zonage géo)
    lettre_fige = get_column_letter(col_geo_max + 1)
    worksheet.freeze_panes = f"{lettre_fige}3"

def _normaliser_nom_colonne(nom_colonne):
    """Normalise un libellé de colonne pour comparer des variantes de nom."""
    if nom_colonne is None:
        return ""
    texte = str(nom_colonne).strip().casefold()
    texte = unicodedata.normalize("NFKD", texte)
    texte = ''.join(ch for ch in texte if not unicodedata.combining(ch))
    texte = re.sub(r"[^a-z0-9]+", " ", texte)
    return re.sub(r"\s+", " ", texte).strip()


def _retrouver_entete_modele(nom_colonne, entetes_modele):
    """Cherche un libellé de colonne compatible dans le modèle, y compris avec des variations de libellés."""
    if nom_colonne in entetes_modele:
        return nom_colonne
    nom_normalise = _normaliser_nom_colonne(nom_colonne)
    for entete in entetes_modele:
        if entete is None:
            continue
        if _normaliser_nom_colonne(entete) == nom_normalise:
            return entete
    return None


def preparer_dataframe_modele(nom_onglet: str, dataframe: pd.DataFrame) -> pd.DataFrame:
    """Adapte les colonnes calculées aux libellés du modèle Excel BDDMAS."""
    resultat = dataframe.drop(columns=COLONNES_ZONAGE_RETIREES, errors="ignore")
    if nom_onglet.startswith("PETITE ENFANCE"):
        resultat = resultat.rename(columns=ZONAGE_PETITE_ENFANCE)
        resultat = resultat.drop(columns=["Superficie en km²"], errors="ignore")
        colonnes_ages = [f"Age : {age}" for age in range(18) if f"Age : {age}" in resultat.columns]
        resultat["TOTAL 0-17"] = resultat[colonnes_ages].sum(axis=1)
    if nom_onglet.startswith("COU FAM"):
        resultat = resultat.rename(columns={
            "Pop 15-24 ans en 2022 (princ)": "Pop 15-24 ans",
            "Pop 25-54 ans en 2022 (princ)": "Pop 25-54 ans",
        })
    if nom_onglet.startswith("CNAM"):
        resultat = resultat.rename(columns=RENOMMAGES_CNAM)
        resultat = resultat.rename(columns={
            colonne: re.sub(r" \(ex-(?:CMUC|ACS)\)", "", colonne)
            for colonne in resultat.columns
            if isinstance(colonne, str) and "(ex-CMUC)" in colonne or isinstance(colonne, str) and "(ex-ACS)" in colonne
        })
    colonnes_techniques = {
        "iris",
        "com",
        "typ iris",
        "lab iris",
        "commune ou arm",
        "type d iris",
        "label de l iris",
        "code geographique",
        "code commune",
        "libelle geographique",
        "libelle de la commune",
    }
    resultat = resultat.drop(
        columns=[
            colonne
            for colonne in resultat.columns
            if _normaliser_nom_colonne(colonne) in colonnes_techniques
        ],
        errors="ignore",
    )
    resultat = resultat.drop(
        columns=[
            colonne
            for colonne in resultat.columns
            if _normaliser_nom_colonne(colonne) in {"note", "note de diffusion"}
        ],
        errors="ignore",
    )
    return resultat


def creer_modele_si_absent(dico_fichiers: dict, chemin_modele: str, chemin_infos_onglets: str) -> None:
    """Crée automatiquement un classeur modèle vide conforme au gabarit métier."""
    if not chemin_modele or os.path.exists(chemin_modele):
        return

    titres = charger_titres_onglets(chemin_infos_onglets)
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)

    feuille_accueil = workbook.create_sheet("Présentation")
    feuille_accueil.sheet_view.showGridLines = False
    feuille_accueil.merge_cells("A1:C1")
    feuille_accueil["A1"] = "Base MAS - INSEE mise à jour 2022"
    feuille_accueil["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    feuille_accueil["A1"].fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    feuille_accueil["A1"].alignment = Alignment(horizontal="center", vertical="center")
    feuille_accueil.row_dimensions[1].height = 40

    workbook.create_sheet("Detail calcul")

    ordre_themes = [
        "POP", "ACT RES", "COU FAM", "DIP FORMATION", "LOG",
        "PETITE ENFANCE", "CAF", "CNAM", "DEFM", "RPLS",
    ]
    ordre_onglets = [
        nom
        for theme in ordre_themes
        for nom in dico_fichiers
        if nom != "PETITE ENFANCE MILLÉSIME" and re.sub(r"\s+\d{4}$", "", nom) == theme
    ]

    for nom_onglet in ordre_onglets:
        if nom_onglet not in dico_fichiers:
            continue
        dataframe = dico_fichiers[nom_onglet]
        feuille = workbook.create_sheet(title=nom_onglet)
        donnees = preparer_dataframe_modele(nom_onglet, dataframe)
        titre = titre_onglet(titres, nom_onglet)
        feuille.cell(row=1, column=1, value=titre)

        if nom_onglet.startswith("PETITE ENFANCE"):
            colonnes_model = [ZONAGE_PETITE_ENFANCE.get(nom, nom) for nom in COLONNES_ZONAGE]
        else:
            colonnes_model = list(COLONNES_ZONAGE)

        colonnes_fin = [col for col in donnees.columns if col not in colonnes_model]
        colonnes_totales = colonnes_model + colonnes_fin
        for index, nom_colonne in enumerate(colonnes_totales, start=1):
            feuille.cell(row=2, column=index, value=nom_colonne)

    for nom_onglet in dico_fichiers:
        if nom_onglet in ordre_onglets or nom_onglet in {"Présentation", "Detail calcul", "PETITE ENFANCE MILLÉSIME"}:
            continue
        dataframe = dico_fichiers[nom_onglet]
        feuille = workbook.create_sheet(title=nom_onglet)
        donnees = preparer_dataframe_modele(nom_onglet, dataframe)
        titre = titre_onglet(titres, nom_onglet)
        feuille.cell(row=1, column=1, value=titre)
        for index, nom_colonne in enumerate(donnees.columns, start=1):
            feuille.cell(row=2, column=index, value=nom_colonne)

    os.makedirs(os.path.dirname(os.path.abspath(chemin_modele)), exist_ok=True)
    workbook.save(chemin_modele)


def _generer_depuis_modele(
    dico_fichiers: dict,
    chemin_sortie: str,
    chemin_infos_onglets: str,
    chemin_modele: str,
    chemin_calculs: str | None = None,
) -> None:
    """Remplit une copie du classeur modèle sans conserver ses lignes d'exemple."""
    titres = charger_titres_onglets(chemin_infos_onglets)
    workbook = load_workbook(chemin_modele)
    feuilles_fixes = {"Présentation", "Detail calcul"}

    for worksheet in workbook.worksheets:
        if worksheet.title in feuilles_fixes:
            continue
        for row in worksheet.iter_rows(min_row=3, max_row=worksheet.max_row):
            for cell in row:
                cell.value = None

    for nom_onglet, dataframe in dico_fichiers.items():
        if nom_onglet == "PETITE ENFANCE MILLÉSIME":
            continue
        if nom_onglet not in workbook.sheetnames:
            theme = nom_onglet.rsplit(" ", 1)[0] if " " in nom_onglet else nom_onglet
            feuilles_theme = [
                titre
                for titre in workbook.sheetnames
                if " " in titre and titre.rsplit(" ", 1)[0] == theme
            ]
            tables_theme = [
                nom
                for nom in dico_fichiers
                if " " in nom and nom.rsplit(" ", 1)[0] == theme and nom != "PETITE ENFANCE MILLÉSIME"
            ]
            if len(feuilles_theme) == 1 and len(tables_theme) == 1:
                workbook[feuilles_theme[0]].title = nom_onglet[:31]

        if nom_onglet not in workbook.sheetnames:
            if nom_onglet == "PETITE ENFANCE MILLÉSIME" and "PETITE ENFANCE 2022" in workbook.sheetnames:
                worksheet = workbook.copy_worksheet(workbook["PETITE ENFANCE 2022"])
                worksheet.title = nom_onglet
            else:
                worksheet = workbook.create_sheet(nom_onglet[:31])
                worksheet.cell(row=2, column=1, value="IRIS")
        else:
            worksheet = workbook[nom_onglet]

        donnees = preparer_dataframe_modele(nom_onglet, dataframe)
        entetes_modele = [
            worksheet.cell(row=2, column=colonne).value
            for colonne in range(1, worksheet.max_column + 1)
        ]

        if nom_onglet.startswith("DEFM"):
            for colonne, entete in enumerate(entetes_modele, start=1):
                entete_normalise = _normaliser_nom_colonne(entete)
                if entete_normalise in ALIAS_ENTETES_DEFM:
                    worksheet.cell(row=2, column=colonne, value=ALIAS_ENTETES_DEFM[entete_normalise])
            entetes_modele = [
                worksheet.cell(row=2, column=colonne).value
                for colonne in range(1, worksheet.max_column + 1)
            ]

        colonnes_obsoletes = [
            colonne
            for colonne, entete in enumerate(entetes_modele, start=1)
            if entete is not None and _retrouver_entete_modele(entete, donnees.columns) is None
        ]
        for colonne in reversed(colonnes_obsoletes):
            worksheet.delete_cols(colonne)
        # Colonnes grises remises dans l'ordre de COLONNES_ZONAGE (mêmes styles : seuls les en-têtes bougent)
        noms_zonage = [
            ZONAGE_PETITE_ENFANCE.get(nom, nom) if nom_onglet.startswith("PETITE ENFANCE") else nom
            for nom in COLONNES_ZONAGE
        ]
        presents = [nom for nom in noms_zonage if nom in donnees.columns]
        positions = sorted(
            colonne
            for colonne in range(1, worksheet.max_column + 1)
            if worksheet.cell(row=2, column=colonne).value in presents
        )
        if len(positions) == len(presents):
            for colonne, nom in zip(positions, presents):
                worksheet.cell(row=2, column=colonne, value=nom)
        entetes_modele = [
            worksheet.cell(row=2, column=colonne).value
            for colonne in range(1, worksheet.max_column + 1)
        ]
        derniere_col = max(
            (c for c, e in enumerate(entetes_modele, start=1) if e is not None),
            default=0,
        )
        
        if worksheet.max_column > derniere_col:
            worksheet.delete_cols(derniere_col + 1, worksheet.max_column - derniere_col)
        entetes_modele = entetes_modele[:derniere_col]

        colonnes_nouvelles = [
            col for col in donnees.columns
            if _retrouver_entete_modele(col, entetes_modele) is None
        ]
        
        position_insertion = next(
            (
                c for c, e in enumerate(entetes_modele, start=1)
                if _normaliser_nom_colonne(e) == _normaliser_nom_colonne("Superficie en km²")
            ),
            len(entetes_modele) + 1,
        )
        if colonnes_nouvelles:
            if position_insertion <= worksheet.max_column:
                worksheet.insert_cols(position_insertion, len(colonnes_nouvelles))
            colonne_modele = max(position_insertion - 1, 1)
            largeur = worksheet.column_dimensions[get_column_letter(colonne_modele)].width
            for decalage, nom_colonne in enumerate(colonnes_nouvelles):
                cible = position_insertion + decalage
                for ligne in range(1, worksheet.max_row + 1):
                    source = worksheet.cell(row=ligne, column=colonne_modele)
                    if source.has_style:
                        worksheet.cell(row=ligne, column=cible)._style = copy(source._style)
                worksheet.cell(row=2, column=cible, value=nom_colonne)
                if largeur:
                    worksheet.column_dimensions[get_column_letter(cible)].width = largeur
            entetes_modele = [
                worksheet.cell(row=2, column=colonne).value
                for colonne in range(1, worksheet.max_column + 1)
            ]

        derniere_ligne = len(donnees) + 2

        if derniere_ligne > worksheet.max_row and worksheet.max_row >= 3:
            for ligne in range(worksheet.max_row + 1, derniere_ligne + 1):
                for colonne in range(1, worksheet.max_column + 1):
                    cellule = worksheet.cell(row=ligne, column=colonne)
                    modele_cellule = worksheet.cell(row=3, column=colonne)
                    if modele_cellule.has_style:
                        cellule._style = copy(modele_cellule._style)
                    if modele_cellule.number_format:
                        cellule.number_format = modele_cellule.number_format

        lignes = tqdm(
            donnees.iterrows(),
            total=len(donnees),
            desc=f"Écriture Excel : {nom_onglet}",
            unit="lignes",
            leave=False,
            dynamic_ncols=True,
        )
        for ligne, (_, enregistrement) in enumerate(lignes, start=3):
            for colonne, entete in enumerate(entetes_modele, start=1):
                nom_colonne_source = None
                for colonne_df in donnees.columns:
                    if colonne_df == entete or _normaliser_nom_colonne(colonne_df) == _normaliser_nom_colonne(entete):
                        nom_colonne_source = colonne_df
                        break
                if nom_colonne_source is None:
                    continue
                valeur = enregistrement[nom_colonne_source]
                if pd.isna(valeur):
                    valeur = None
                elif hasattr(valeur, "item"):
                    valeur = valeur.item()
                worksheet.cell(row=ligne, column=colonne, value=valeur)

        titre = titre = titre_onglet(titres, nom_onglet)
        worksheet.cell(row=1, column=1, value=titre)
        # Colonnes grises = jusqu'au libellé IRIS inclus
        nombre_colonnes_geo = (
            entetes_modele.index("Libellé de l'Iris") + 1
            if "Libellé de l'Iris" in entetes_modele
            else len(COLONNES_ZONAGE)
        )
        worksheet.cell(row=1, column=nombre_colonnes_geo, value="Total")
        derniere_ligne_donnees = max(derniere_ligne, 3)
        for colonne in range(nombre_colonnes_geo + 1, worksheet.max_column + 1):
            lettre = get_column_letter(colonne)
            worksheet.cell(
                row=1,
                column=colonne,
                value=f"=SUBTOTAL(9,{lettre}3:{lettre}{derniere_ligne_donnees})",
            )

    for worksheet in workbook.worksheets:
        for table in worksheet.tables.values():
            min_col, min_row, max_col, max_row = range_boundaries(table.ref)
            max_col = max(
                (cell.column for cell in worksheet[min_row] if cell.value is not None),
                default=min_col,
            )
            max_row = max(worksheet.max_row, min_row + 1)
            largeur = max_col - min_col + 1
            table_columns = list(table.tableColumns[:largeur])
            while len(table_columns) < largeur:
                index = len(table_columns) + 1
                table_columns.append(TableColumn(id=index, name=f"Column{index}"))

            noms_utilises = set()
            for index, table_column in enumerate(table_columns):
                colonne = min_col + index
                entete = worksheet.cell(row=min_row, column=colonne).value
                nom_entete = str(entete) if entete is not None else f"Column{index + 1}"
                nom_unique = nom_entete
                suffixe = 2
                while nom_unique in noms_utilises:
                    nom_unique = f"{nom_entete}_{suffixe}"
                    suffixe += 1
                if nom_unique != entete:
                    worksheet.cell(row=min_row, column=colonne, value=nom_unique)
                noms_utilises.add(nom_unique)
                table_column.id = index + 1
                table_column.name = nom_unique

            table.tableColumns = table_columns
            table.ref = (
                f"{get_column_letter(min_col)}{min_row}:"
                f"{get_column_letter(max_col)}{max_row}"
            )
            if table.autoFilter is not None:
                table.autoFilter.ref = table.ref

    mettre_a_jour_pages_info(workbook, dico_fichiers, chemin_calculs)
    appliquer_mise_en_page(workbook)
    ajouter_portrait_roubaix(workbook)
    os.makedirs(os.path.dirname(os.path.abspath(chemin_sortie)), exist_ok=True)
    workbook.save(chemin_sortie)


def generer_fichier_bddmas(
    dico_fichiers: dict,
    chemin_sortie: str,
    chemin_infos_onglets: str,
    chemin_modele: str = None,
    chemin_calculs: str | None = None,
):
    """
    Génère le classeur Excel final à partir du dictionnaire de DataFrames.
    Si le modèle est absent, il est créé automatiquement avant d'être alimenté.
    """
    if chemin_modele and not os.path.exists(chemin_modele):
        print("MODELE ABSENT : création automatique du fichier de référence...")
        creer_modele_si_absent(dico_fichiers, chemin_modele, chemin_infos_onglets)

    if chemin_modele and os.path.exists(chemin_modele):
        print("GÉNÉRATION DU CLASSEUR À PARTIR DU MODÈLE...")
        _generer_depuis_modele(dico_fichiers, chemin_sortie, chemin_infos_onglets, chemin_modele, chemin_calculs)
        print(f"Fichier Excel enregistré avec succès sous : {chemin_sortie}")
        return

    print("\nGÉNÉRATION DU FICHIER EXCEL FINAL...")
    titres_onglets = charger_titres_onglets(chemin_infos_onglets)
    
    with pd.ExcelWriter(chemin_sortie, engine='openpyxl') as writer:
        
        # --- 1. Création de l'onglet de Présentation ---
        ws_accueil = writer.book.create_sheet("Présentation", 0)
        ws_accueil.sheet_view.showGridLines = False 
        
        ws_accueil.merge_cells('A1:C1')
        ws_accueil['A1'].value = "Base MAS - INSEE mise à jour 2022"
        ws_accueil['A1'].font = Font(bold=True, size=16, color="FFFFFF")
        ws_accueil['A1'].fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
        ws_accueil['A1'].alignment = Alignment(horizontal="center", vertical="center")
        ws_accueil.row_dimensions[1].height = 40
        
        # --- 2. Boucle sur les onglets thématiques ---
        for nom_onglet, df in dico_fichiers.items():
            print(f" -> Export et formatage de l'onglet : {nom_onglet}")
            
            df = df.drop(columns=COLONNES_ZONAGE_RETIREES, errors="ignore")
            # Export du DataFrame avec startrow=1 (laisse la ligne 1 vierge pour nos Totaux)
            df.to_excel(writer, sheet_name=nom_onglet, index=False, startrow=1)
            
            # Récupération de la feuille pour le formatage
            worksheet = writer.sheets[nom_onglet]
            titre_principal = titre_onglet(titres_onglets, nom_onglet)
            
            formater_onglet_thematique(worksheet, df, titre_principal)
            
    print(f"Fichier Excel enregistré avec succès sous : {chemin_sortie}")