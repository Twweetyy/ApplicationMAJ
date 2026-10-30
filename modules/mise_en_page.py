"""Mise en page des onglets de données : formats de nombres, largeurs, volets figés."""
import re
from copy import copy

from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import get_column_letter, range_boundaries

from .colonnes_auto import regle_colonne
from .indicateurs import _precision_arrondi

FEUILLES_FIXES = {"Présentation", "Detail calcul"}
FORMATS = {0: "#,##0", 1: "#,##0.0", 2: "#,##0.00"}
LARGEUR_MIN, LARGEUR_MAX = 11, 32
LARGEUR_TEXTE_MAX = 45
# Colonnes grises (« Mairies de quartiers », « Libellé de l'Iris ») : bornes de largeur
LARGEUR_GEO_MIN, LARGEUR_GEO_MAX = 12, 30
LIGNES_ENTETE_MAX = 5
LIGNES_TEXTE_MAX = 2
# Largeur Excel consommée par caractère, et place prise par le bouton de filtre
LARGEUR_CARACTERE, MARGE_FILTRE = 1.1, 2
# En-tête orange = indicateur (taux, part, médiane...) : pas de total
REMPLISSAGE_INDICATEUR = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
# Thèmes dont aucune colonne ne s'additionne (quartiles, médianes, taux Filosofi)
THEMES_NON_ADDITIFS = ("REV DEC", "REV DISP")
# Colonnes du modèle sans source de données : supprimées si vides
COLONNES_SANS_SOURCE = {"Crèches"}


def _est_additive(nom_colonne: str, nom_feuille: str, precision: int) -> bool:
    """Une somme sur Roubaix n'a de sens que pour des effectifs (et la superficie)."""
    if nom_feuille.startswith(THEMES_NON_ADDITIFS):
        return False
    # Colonne créée par modules/colonnes_auto.py : additivité donnée par le module
    regle = regle_colonne(nom_colonne, nom_feuille)
    if regle and regle.get("additive") is not None:
        return regle["additive"]
    nom = nom_colonne.strip().casefold()
    if nom.startswith("superficie"):
        return True
    if precision >= 1:
        return False
    return not re.search(r"moyen|médian|median|quartile|décile|decile|loyer|€", nom)


def _nb_lignes(texte: str, largeur: float) -> int:
    """Nombre de lignes d'un texte renvoyé à la ligne dans une colonne de cette largeur."""
    capacite = max((largeur - MARGE_FILTRE) / LARGEUR_CARACTERE, 1)
    lignes, longueur = 1, 0
    for mot in texte.split():
        if longueur and longueur + 1 + len(mot) > capacite:
            lignes += 1
            longueur = len(mot)
        else:
            longueur += len(mot) + (1 if longueur else 0)
    return lignes


def _largeur_pour(texte: str, largeur: int, lignes_max: int, largeur_max: int) -> int:
    """Élargit la colonne jusqu'à ce que le texte tienne sur lignes_max lignes."""
    while largeur < largeur_max and _nb_lignes(texte, largeur) > lignes_max:
        largeur += 1
    return largeur


def _est_orange(cellule) -> bool:
    couleur = cellule.fill.fgColor.rgb
    return isinstance(couleur, str) and couleur[-6:].upper() == "F8CBAD"


def _regrouper_colonnes(ws, premiere_colonne: int) -> None:
    """Colonnes bleues (effectifs) à gauche, orange (indicateurs) à droite, ordre conservé."""
    colonnes = list(range(premiere_colonne, ws.max_column + 1))
    ordre = sorted(colonnes, key=lambda c: _est_orange(ws.cell(row=2, column=c)))
    if ordre == colonnes:
        return
    derniere_ligne = ws.max_row
    contenu = {
        c: [(ws.cell(row=l, column=c).value, copy(ws.cell(row=l, column=c)._style)) for l in range(1, derniere_ligne + 1)]
        for c in colonnes
    }
    largeurs = {c: ws.column_dimensions[get_column_letter(c)].width for c in colonnes}
    for nouvelle, ancienne in zip(colonnes, ordre):
        lettre = get_column_letter(nouvelle)
        for ligne, (valeur, style) in enumerate(contenu[ancienne], start=1):
            cellule = ws.cell(row=ligne, column=nouvelle)
            cellule.value = valeur
            cellule._style = style
        ws.column_dimensions[lettre].width = largeurs[ancienne]
        total = ws.cell(row=1, column=nouvelle)
        if isinstance(total.value, str) and total.value.startswith("=SUBTOTAL"):
            total.value = re.sub(r"[A-Z]+(\d+):[A-Z]+(\d+)", rf"{lettre}\1:{lettre}\2", total.value)
    # Tableau Excel : ses colonnes doivent suivre le même ordre que les en-têtes
    for tableau in ws.tables.values():
        min_col, _, max_col, _ = range_boundaries(tableau.ref)
        if not (min_col <= premiere_colonne and colonnes[-1] <= max_col):
            continue
        anciennes = list(tableau.tableColumns)
        debut = premiere_colonne - min_col
        tableau.tableColumns = anciennes[:debut] + [anciennes[c - min_col] for c in ordre] + anciennes[debut + len(ordre):]


def _derniere_ligne(ws) -> int:
    ligne = ws.max_row
    while ligne > 3 and all(cellule.value in (None, "") for cellule in ws[ligne]):
        ligne -= 1
    return ligne


def _supprimer_colonnes_vides(ws, noms=COLONNES_SANS_SOURCE) -> None:
    """Retire les colonnes du modèle sans source (ex. « Crèches ») quand elles sont vides."""
    derniere_ligne = _derniere_ligne(ws)
    for colonne in range(ws.max_column, 0, -1):
        if ws.cell(row=2, column=colonne).value not in noms:
            continue
        if any(ws.cell(row=ligne, column=colonne).value not in (None, "") for ligne in range(3, derniere_ligne + 1)):
            continue
        # Largeurs : openpyxl ne les décale pas avec delete_cols
        largeurs = {
            c: ws.column_dimensions[get_column_letter(c)].width
            for c in range(colonne + 1, ws.max_column + 1)
        }
        ws.delete_cols(colonne)
        for c, largeur in largeurs.items():
            if largeur:
                ws.column_dimensions[get_column_letter(c - 1)].width = largeur
        # Tableau Excel : plage et colonnes
        for tableau in ws.tables.values():
            min_col, min_row, max_col, max_row = range_boundaries(tableau.ref)
            if not min_col <= colonne <= max_col:
                continue
            tableau.ref = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col - 1)}{max_row}"
            del tableau.tableColumns[colonne - min_col]
            for i, colonne_tableau in enumerate(tableau.tableColumns, start=1):
                colonne_tableau.id = i
            if tableau.autoFilter is not None:
                tableau.autoFilter.ref = tableau.ref
        # Ligne Total : les SUBTOTAL gardent l'ancienne lettre, on les réécrit
        for c in range(colonne, ws.max_column + 1):
            cellule = ws.cell(row=1, column=c)
            if isinstance(cellule.value, str) and cellule.value.startswith("=SUBTOTAL"):
                lettre = get_column_letter(c)
                cellule.value = f"=SUBTOTAL(9,{lettre}3:{lettre}{derniere_ligne})"


def _mettre_en_page_onglet(ws) -> None:
    _supprimer_colonnes_vides(ws)
    entetes = [cellule.value for cellule in ws[2]]
    # Colonnes grises = jusqu'au libellé IRIS inclus (Mairies de quartiers + Libellé de l'Iris)
    nombre_colonnes_geo = entetes.index("Libellé de l'Iris") + 1 if "Libellé de l'Iris" in entetes else 2
    derniere_ligne = _derniere_ligne(ws)
    # Le modèle petite enfance s'ouvre en F1 : sans ça, A-E restent cachées derrière le volet
    ws.sheet_view.topLeftCell = "A1"
    ws.freeze_panes = ws.cell(row=3, column=nombre_colonnes_geo + 1)

    # Colonnes grises : largeur ajustée au contenu (delete_cols ne décale pas les largeurs du modèle)
    for colonne in range(1, nombre_colonnes_geo + 1):
        textes = [str(ws.cell(row=ligne, column=colonne).value or "") for ligne in range(3, derniere_ligne + 1)]
        plus_long = max((len(t) for t in textes), default=0)
        mot_entete = max((len(mot) for mot in str(entetes[colonne - 1] or "").split()), default=0)
        largeur = min(max(plus_long + 3, mot_entete + MARGE_FILTRE + 2, LARGEUR_GEO_MIN), LARGEUR_GEO_MAX)
        ws.column_dimensions[get_column_letter(colonne)].width = largeur
    # Lignes d'IRIS : hauteur automatique (le modèle gardait celles des anciens libellés longs)
    for ligne in range(3, derniere_ligne + 1):
        ws.row_dimensions[ligne].height = None
    # Titre : fusionné sur les colonnes grises (A1 ne déborde plus, la colonne B étant occupée)
    titre = ws.cell(row=1, column=1)
    if isinstance(titre.value, str) and titre.value and nombre_colonnes_geo > 1:
        for colonne in range(2, nombre_colonnes_geo + 1):
            ws.cell(row=1, column=colonne).value = None
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=nombre_colonnes_geo)
        titre.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        largeur_titre = sum(
            ws.column_dimensions[get_column_letter(c)].width for c in range(1, nombre_colonnes_geo + 1)
        )
        ws.row_dimensions[1].height = max(_nb_lignes(titre.value, largeur_titre), 2) * 15 + 2

    lignes_entete = 2
    for colonne in range(nombre_colonnes_geo + 1, len(entetes) + 1):
        entete = entetes[colonne - 1]
        if entete in (None, ""):
            continue
        entete = str(entete)
        lettre = get_column_letter(colonne)

        # Colonne de texte (ex. « Incertitude » REV DEC) : large, 2 lignes max par IRIS
        textes = [
            str(ws.cell(row=ligne, column=colonne).value)
            for ligne in range(3, derniere_ligne + 1)
            if isinstance(ws.cell(row=ligne, column=colonne).value, str)
        ]
        if textes:
            plus_long_texte = max(textes, key=len)
            largeur = _largeur_pour(plus_long_texte, LARGEUR_MIN, LIGNES_TEXTE_MAX, LARGEUR_TEXTE_MAX)
            largeur = _largeur_pour(entete, largeur, LIGNES_ENTETE_MAX, LARGEUR_TEXTE_MAX)
            ws.column_dimensions[lettre].width = largeur
            lignes_entete = max(lignes_entete, _nb_lignes(entete, largeur))
            total = ws.cell(row=1, column=colonne)
            if isinstance(total.value, str) and total.value.startswith("=SUBTOTAL"):
                total.value = None
            continue

        precision = _precision_arrondi(entete, ws.title)
        if precision is None:
            continue
        cellules = [
            ws.cell(row=ligne, column=colonne)
            for ligne in range(3, derniere_ligne + 1)
            if isinstance(ws.cell(row=ligne, column=colonne).value, (int, float))
            and not isinstance(ws.cell(row=ligne, column=colonne).value, bool)
        ]
        # Pas plus de décimales affichées que la source n'en contient (RPLS : 9,3 et non 9,30)
        decimales_source = max(
            (len(f"{cellule.value:.{precision}f}".rstrip("0").split(".")[1]) if precision else 0 for cellule in cellules),
            default=precision,
        )
        precision = min(precision, max(decimales_source, 1 if precision else 0))
        format_nombre = FORMATS.get(precision, "#,##0.00")

        longueur_valeurs = 0
        for cellule in cellules:
            cellule.number_format = format_nombre
            longueur_valeurs = max(longueur_valeurs, len(f"{cellule.value:,.{precision}f}"))

        additive = _est_additive(entete, ws.title, precision)
        total = ws.cell(row=1, column=colonne)
        if isinstance(total.value, str) and total.value.startswith("=SUBTOTAL"):
            if additive:
                total.number_format = format_nombre
            else:
                total.value = None
        if not additive:
            ws.cell(row=2, column=colonne).fill = REMPLISSAGE_INDICATEUR

        # Largeur : l'en-tête doit tenir sur 5 lignes, les nombres sur une seule
        plus_long_mot = max((len(mot) for mot in entete.split()), default=0)
        largeur = min(max(plus_long_mot + MARGE_FILTRE + 1, longueur_valeurs + 3, LARGEUR_MIN), LARGEUR_MAX)
        largeur = _largeur_pour(entete, largeur, LIGNES_ENTETE_MAX, LARGEUR_MAX)
        ws.column_dimensions[lettre].width = largeur
        lignes_entete = max(lignes_entete, _nb_lignes(entete, largeur))

    ws.row_dimensions[2].height = lignes_entete * 15 + 5
    _regrouper_colonnes(ws, nombre_colonnes_geo + 1)


def appliquer_mise_en_page(workbook) -> None:
    """Formats de nombres, largeurs de colonnes et volets figés sur chaque onglet de données."""
    for ws in workbook.worksheets:
        if ws.title in FEUILLES_FIXES or ws.max_row < 3:
            continue
        _mettre_en_page_onglet(ws)
