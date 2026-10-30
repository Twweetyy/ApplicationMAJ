"""Classeur de la Base permanente des équipements (BPE) de l'Insee, pour Roubaix ou la MEL.

Entrée : le fichier BPE de l'Insee (national, départemental ou déjà filtré), en CSV ou Parquet,
et le fichier de métadonnées varmod (BPE<AA>_anonymisee_varmod.csv). Sans varmod fourni,
celui de dictionnaires/bpe/ est utilisé.

Sortie : un classeur Excel avec les onglets Accueil, BPE, Recap (comptages par IRIS ou par
commune), un onglet par domaine d'équipement et le dictionnaire Insee.

Périmètres :
  roubaix : équipements de Roubaix, avec libellé IRIS, secteur mairie, quartier historique
            et grand quartier ; Recap par IRIS ;
  mel     : équipements des 95 communes de la MEL ; Recap par commune.
"""
import csv
import os

import pandas as pd
import yaml
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .lecture import est_parquet, lire_par_blocs

CODE_ROUBAIX = "59512"
PERIMETRES = {"roubaix": "Roubaix", "mel": "MEL"}
TAILLE_BLOC = 200_000
LIEN_INSEE = "https://www.insee.fr/fr/statistiques/8217525"
FICHIER_VARMOD_DEFAUT = os.path.join("bpe", "BPE25_anonymisee_varmod.csv")

# Code Insee -> libellé de l'en-tête dans le classeur
LIBELLES_COLONNES = {
    "AN": "Année",
    "NOMRS": "Nom Raison sociale de l'équipement",
    "CNOMRS": "Complément du nom ",
    "CADR": "Complément d'adresse",
    "TYPEQU": "Code Type d'équipement",
    "SDOM": "Code  Sous Domaine d'appartenance",
    "DOM": "Code Domaine d'appartenance",
    "NUMVOIE": "Numéro de voie de l'adresse",
    "INDREP": "Indice de Repetition ",
    "TYPVOIE": "Type de la voie",
    "LIBVOIE": "Libellé de la voie ",
    "DEPCOM": "Code Insee",
    "LIBCOM": "Commune",
    "STATUT_DIFFUSION": "Statut de diffusion Sirene",
    "CANTINE": "Présence ou absence d'une cantine",
    "INTERNAT": "Présence ou absence d’un internat",
    "RPI": "Type de regroupement pédagogique intercommunal",
    "EP": "Présence ou absence d'un dispositif d'éducation prioritaire",
    "CL_PGE": "Présence ou absence d'une classe préparatoire aux grandes écoles en lycée",
    "SECT": "Appartenance au secteur public ou privé d'enseignement",
    "SECTEUR": "Appartenance au secteur public ou privé d'enseignement",
    "ACCES_AIRE_PRATIQUE": "Présence d’un aménagement permettant l’accessibilité d’au moins une aire de pratique aux personnes handicapées",
    "ACCES_LIBRE": "Présence d’un équipement en accès libre permanent",
    "ACCES_SANITAIRE": "Présence d’un aménagement permettant l’accessibilité d’au moins un sanitaire, aux personnes handicapées ",
    "ACCES_VESTIAIRE": "Présence d’un aménagement permettant l’accessibilité d’au moins un vestiaire et une douche  aux personnes handicapées moteur",
    "CATEGORIE": "Catégorie d’établissement recevant du public (ERP)",
    "PRES_DOUCHE": "Présencede douches",
    "PRES_SANITAIRE": "Présence de sanitaires",
    "SAISONNIER": "Equipement à l’ouverture exclusivement saisonnière",
    "COUVERT": "Présence ou absence d’équipement(s) avec au moins une partie couverte",
    "ECLAIRE": "Présence ou absence d'équipement(s) avec au moins une partie éclairée",
    "TYPE": "Catégorie d'équipement",
    "MULTIPLEXE": "Présence d’un cinéma multiplexe",
    "STRUCTURE_EXERCICE": "Structure d'exercice de l'activité",
    "SPECIALITE": "Spécialité du professionnel",
    "ACCUEIL": "Type d'accueil de l'infrastructure",
    "ITINERANCE": "Présence d’une structure itinérante",
    "MODE_GESTION": "Mode de gestion de l'infrastructure",
    "SSTYPHEB": "Sous-type d'hébergement",
    "TYPERESTO": "Type de lieux",
    "IMPLANTATION_STATION": "Implantation d’une station de recharge",
    "CAPACITE_D_ACCUEIL": "Capacité de l'équipement",
    "INDIC_CAPA": "Indicateur de concernement des valeurs de capacité",
    "NBEQUIDENT": "Nombre d'aires de pratique identiques dans l'équipement",
    "INDIC_NBEQUIDENT": "Indicateur de concernement des valeurs du nombre d’infrastructures au sein de l’équipement",
    "NBSALLES": "Nombre de salles",
    "INDIC_NBSALLES": "Indicateur de concernement des valeurs du nombre de salles",
    "NBLIEUX": "Nombre de lieux",
    "INDIC_NBLIEUX": "Indicateur de concernement des valeurs du nombre de lieux",
    "NB_PDC": "Nombre de points de charge dans une station de recharge",
    "INDIC_NB_PDC": "Indicateur de concernement des valeurs du nombre de points de charge dans une station de recharge",
    "NB_PDC_PA": "Nombre de points de recharge à paiement à l’acte dans une station de recharge",
    "INDIC_NB_PDC_PA": "Indicateur de concernement des valeurs du nombre de points de recharge à paiement à l’acte",
    "NB_PDC_ACCELEREE": "Nombre de points de recharge accélérée dans une station de recharge",
    "INDIC_NB_PDC_ACCELEREE": "Indicateur de concernement des valeurs du nombre de points de recharge accélérée",
    "NB_PDC_LENTE": "Nombre de points de recharge lente dans une station de recharge",
    "INDIC_NB_PDC_LENTE": "Indicateur de concernement des valeurs du nombre de points de recharge lente",
    "NB_PDC_RAPIDE": "Nombre de points de recharge rapide dans une station de recharge",
    "INDIC_NB_PDC_RAPIDE": "Indicateur de concernement des valeurs du nombre de points de recharge rapide",
    "NB_PDC_ULTRARAPIDE": "Nombre de points de recharge ultra-rapide dans une station de recharge",
    "INDIC_NB_PDC_ULTRARAPIDE": "Indicateur de concernement des valeurs du nombre de points de recharge ultra-rapide",
    "NB_JOURS_OUVERT": "Nombre de jours d’ouverture hebdomadaire d’une station de recharge",
    "INDIC_NB_JOURS_OUVERT": "Indicateur de concernement des valeurs du nombre de jours d’ouverture hebdomadaire d’une station de recharge",
    "QUALITE_GEOLOC": "Indicateur de qualité de géolocalisation",
    "EPSG": "Code EPSG du système de coordonnées",
    "QUALITE_XY": "Indicateur de précision du repérage par l’adresse",
    "TR_DIST_PRECISION": "Erreur maximum de positionnement dans la voie",
    "DCIRIS": "Code IRIS",
    "QUALI_IRIS": "Indicateur de qualité du géoréférencement dans l'IRIS",
    "QP2024": "Quartier prioritaire de la politique dela ville 2024",
    "QUALI_QP2024": "Indicateur de qualité du géoréférencement dans le quartier prioritaire de la politique de la ville 2024",
    "ZUS": "Zone urbaine sensible",
    "QUALI_ZUS": "Indicateur de qualité du géoréférencement dans la zone urbaine sensible",
    "SIRET": "Numero SIRET",
    "LAMBERT_X": "Lambert X",
    "LAMBERT_Y": "lambert Y",
    # Nouveautés 2025
    "CODPOS": "Code Postal",
    "LATITUDE": "Latitude GPS",
    "LONGITUDE": "Longitude GPS",
    "APET": "Code APE (Activité Principale)",
    "DEP": "Département",
    "REG": "Région",
    "EPCI": "Intercommunalité (EPCI)",
    "UU2020": "Unité urbaine 2020",
    "AAV2020": "Aire d'attraction des villes 2020",
    "BV2022": "Bassin de vie 2022",
    "QP2015": "Quartier prioritaire 2015",
    "DENS3": "Densité (3 niveaux)",
    "DENS7": "Densité (7 niveaux)",
    "QVA": "Quartier de veille active (QVA)",
    "IRISEE": "Commune irisée",
    "GPL": "Gaz de pétrole liquéfié",
    "CAPACITE": "Capacité de l'équipement (v2025)",
}

# Colonnes ajoutées pour Roubaix (dictionnaires du dépôt), placées après la zone urbaine sensible
COLONNES_ROUBAIX = ["Secteur Mairie", "Quartier historique", "Grand quartier"]

# Ordre des colonnes de l'onglet BPE (les colonnes absentes du millésime restent vides)
COLONNES_EXPORT = [
    "OBJECTID", "Année", "Nom Raison sociale de l'équipement", "Complément du nom ", "Complément d'adresse",
    "Code Type d'équipement", "Type d'équipement", "Code  Sous Domaine d'appartenance", "Sous Domaine d'appartenance",
    "Code Domaine d'appartenance", "Domaine d'appartenance", "Numéro de voie de l'adresse", "Indice de Repetition ",
    "Type de la voie", "Libellé de la voie ", "Code Insee", "Commune", "Statut de diffusion Sirene",
    "Présence ou absence d'une cantine", "Présence ou absence d’un internat", "Type de regroupement pédagogique intercommunal",
    "Présence ou absence d'un dispositif d'éducation prioritaire", "Présence ou absence d'une classe préparatoire aux grandes écoles en lycée",
    "Appartenance au secteur public ou privé d'enseignement", "Présence d’un aménagement permettant l’accessibilité d’au moins une aire de pratique aux personnes handicapées",
    "Présence d’un équipement en accès libre permanent", "Présence d’un aménagement permettant l’accessibilité d’au moins un sanitaire, aux personnes handicapées ",
    "Présence d’un aménagement permettant l’accessibilité d’au moins un vestiaire et une douche  aux personnes handicapées moteur",
    "Catégorie d’établissement recevant du public (ERP)", "Présencede douches", "Présence de sanitaires",
    "Equipement à l’ouverture exclusivement saisonnière", "Présence ou absence d’équipement(s) avec au moins une partie couverte",
    "Présence ou absence d'équipement(s) avec au moins une partie éclairée", "Catégorie d'équipement", "Présence d’un cinéma multiplexe",
    "Structure d'exercice de l'activité", "Spécialité du professionnel", "Type d'accueil de l'infrastructure",
    "Présence d’une structure itinérante", "Mode de gestion de l'infrastructure", "Sous-type d'hébergement", "Type de lieux",
    "Implantation d’une station de recharge", "Capacité de l'équipement", "Indicateur de concernement des valeurs de capacité",
    "Nombre d'aires de pratique identiques dans l'équipement", "Indicateur de concernement des valeurs du nombre d’infrastructures au sein de l’équipement",
    "Nombre de salles", "Indicateur de concernement des valeurs du nombre de salles",
    "Nombre de lieux", "Indicateur de concernement des valeurs du nombre de lieux",
    "Nombre de points de charge dans une station de recharge", "Indicateur de concernement des valeurs du nombre de points de charge dans une station de recharge",
    "Nombre de points de recharge à paiement à l’acte dans une station de recharge", "Indicateur de concernement des valeurs du nombre de points de recharge à paiement à l’acte",
    "Nombre de points de recharge accélérée dans une station de recharge", "Indicateur de concernement des valeurs du nombre de points de recharge accélérée",
    "Nombre de points de recharge lente dans une station de recharge", "Indicateur de concernement des valeurs du nombre de points de recharge lente",
    "Nombre de points de recharge rapide dans une station de recharge", "Indicateur de concernement des valeurs du nombre de points de recharge rapide",
    "Nombre de points de recharge ultra-rapide dans une station de recharge", "Indicateur de concernement des valeurs du nombre de points de recharge ultra-rapide",
    "Nombre de jours d’ouverture hebdomadaire d’une station de recharge", "Indicateur de concernement des valeurs du nombre de jours d’ouverture hebdomadaire d’une station de recharge",
    "Indicateur de qualité de géolocalisation", "Code EPSG du système de coordonnées", "Indicateur de précision du repérage par l’adresse",
    "Erreur maximum de positionnement dans la voie", "Code IRIS", "Indicateur de qualité du géoréférencement dans l'IRIS",
    "Quartier prioritaire de la politique dela ville 2024", "Indicateur de qualité du géoréférencement dans le quartier prioritaire de la politique de la ville 2024",
    "Zone urbaine sensible", "Indicateur de qualité du géoréférencement dans la zone urbaine sensible",
    *COLONNES_ROUBAIX,
    "Numero SIRET", "Lambert X", "lambert Y",
    "Code Postal", "Latitude GPS", "Longitude GPS", "Latitude et longitude GPS", "Code APE (Activité Principale)", "Département", "Région",
    "Intercommunalité (EPCI)", "Unité urbaine 2020", "Aire d'attraction des villes 2020", "Bassin de vie 2022",
    "Quartier prioritaire 2015", "Densité (3 niveaux)", "Densité (7 niveaux)", "Quartier de veille active (QVA)",
    "Commune irisée", "Gaz de pétrole liquéfié", "Capacité de l'équipement (v2025)",
]

# Description des onglets de domaine sur la page d'accueil (les domaines absents ne sont pas listés)
DESCRIPTIONS_DOMAINES = {
    "A": "Détail des équipements de type 'Services' (Agences bancaires, postes, garages...).",
    "B": "Détail des commerces de détail, grandes surfaces, boulangeries, supermarchés, etc.",
    "C": "Établissements scolaires, de la maternelle au supérieur.",
    "D": "Infrastructures médicales, sociales, médecins, pharmacies, hôpitaux.",
    "E": "Infrastructures de transport, gares, stations-service, bornes de recharge.",
    "F": "Équipements sportifs, bibliothèques, cinémas, musées, lieux de culte.",
    "G": "Hébergements touristiques, campings, offices de tourisme.",
}

BLEU = "4F81BD"
BLEU_FONCE = "1F497D"
BLEU_CLAIR = "DCE6F1"


class ErreurBPE(ValueError):
    """Fichier BPE inutilisable (colonnes manquantes, aucun équipement dans le périmètre…)."""


def _charger_yml(chemin: str) -> dict[str, str]:
    if not os.path.exists(chemin):
        return {}
    with open(chemin, encoding="utf-8") as fichier:
        return {str(cle): str(valeur).strip() for cle, valeur in (yaml.safe_load(fichier) or {}).items()}


def _separateur_csv(chemin: str) -> str:
    with open(chemin, encoding="utf-8-sig", errors="replace") as fichier:
        entete = fichier.readline()
    try:
        return csv.Sniffer().sniff(entete, delimiters=";,\t").delimiter
    except csv.Error:
        return ";"


def lire_varmod(chemin: str) -> pd.DataFrame:
    df_meta = pd.read_csv(chemin, sep=_separateur_csv(chemin), dtype=str, encoding="utf-8-sig")
    manquantes = {"COD_VAR", "COD_MOD", "LIB_MOD"} - set(df_meta.columns)
    if manquantes:
        raise ErreurBPE(f"Fichier varmod sans les colonnes {', '.join(sorted(manquantes))} : {os.path.basename(chemin)}")
    return df_meta


def _en_texte(valeur) -> str | None:
    """Nombre du Parquet écrit comme dans le CSV de l'Insee : 12.0 -> « 12 », 50.69 -> « 50.69 »."""
    if valeur is None or pd.isna(valeur):
        return None
    if isinstance(valeur, float) and valeur.is_integer():
        return str(int(valeur))
    return str(valeur)


def _lire_parquet_perimetre(chemin_bpe: str, codes_communes: set[str]) -> pd.DataFrame:
    """Filtre DEPCOM fait par DuckDB dans le Parquet : seules les lignes du périmètre sont chargées."""
    import duckdb

    source = chemin_bpe.replace("'", "''")
    connexion = duckdb.connect()
    try:
        colonnes = [ligne[0] for ligne in connexion.execute(f"DESCRIBE SELECT * FROM read_parquet('{source}')").fetchall()]
        if "DEPCOM" not in colonnes:
            raise ErreurBPE("Colonne DEPCOM absente : ce fichier n'est pas une table BPE de l'Insee.")
        df = connexion.execute(
            f"SELECT * FROM read_parquet('{source}') WHERE trim(CAST(DEPCOM AS VARCHAR)) IN (SELECT unnest(?))",
            [sorted(codes_communes)],
        ).df()
    except duckdb.Error as erreur:
        raise ErreurBPE(f"Parquet illisible ({os.path.basename(chemin_bpe)}) : {erreur}") from erreur
    finally:
        connexion.close()
    for col in df.columns:
        if not pd.api.types.is_string_dtype(df[col]) and not pd.api.types.is_object_dtype(df[col]):
            df[col] = df[col].astype(object).map(_en_texte)
    return df.astype(object).where(df.notna(), None)


def extraire_perimetre(chemin_bpe: str, codes_communes: set[str]) -> pd.DataFrame:
    """Lignes du fichier BPE dont DEPCOM est dans codes_communes (fichier national accepté, CSV lu par blocs)."""
    if est_parquet(chemin_bpe):
        return _lire_parquet_perimetre(chemin_bpe, codes_communes)
    lecture = lire_par_blocs(chemin_bpe, TAILLE_BLOC, sep=_separateur_csv(chemin_bpe), dtype=str, encoding="utf-8-sig")
    blocs = []
    for bloc in lecture:
        if "DEPCOM" not in bloc.columns:
            raise ErreurBPE("Colonne DEPCOM absente : ce fichier n'est pas une table BPE de l'Insee.")
        bloc = bloc.astype("string")
        blocs.append(bloc[bloc["DEPCOM"].str.strip().isin(codes_communes)])
    if not blocs:
        raise ErreurBPE("Fichier BPE vide.")
    df = pd.concat(blocs, ignore_index=True)
    # Index exporté par un précédent filtrage (to_csv sans index=False)
    df = df.drop(columns=[col for col in df.columns if str(col).startswith("Unnamed")])
    return df.astype(object).where(df.notna(), None)


def _traduire(df_data: pd.DataFrame, df_meta: pd.DataFrame) -> tuple[pd.DataFrame, dict]:
    """Ajoute les libellés des domaines et types, puis remplace les codes par leurs libellés varmod."""
    dictionnaire = {}
    for cod_var, sous_table in df_meta.dropna(subset=["COD_VAR"]).groupby("COD_VAR"):
        valides = sous_table[sous_table["COD_MOD"].notna() & (sous_table["COD_MOD"] != "liste")]
        if not valides.empty:
            dictionnaire[cod_var] = valides.drop_duplicates("COD_MOD").set_index("COD_MOD")["LIB_MOD"].to_dict()

    for code, libelle in (("DOM", "Domaine d'appartenance"), ("SDOM", "Sous Domaine d'appartenance"),
                          ("TYPEQU", "Type d'équipement")):
        if code in df_data.columns:
            df_data[libelle] = df_data[code].map(dictionnaire.get(code, {})).fillna(df_data[code])

    for col in df_data.columns:
        if col in dictionnaire and col not in {"DOM", "SDOM", "TYPEQU"}:
            df_data[col] = df_data[col].map(dictionnaire[col]).fillna(df_data[col])
    # Codes génériques que le varmod ne décrit pas toujours
    return df_data.replace({"_Z": "Sans objet", "_U": "Indéterminé"}), dictionnaire


def preparer_bpe(df_data: pd.DataFrame, df_meta: pd.DataFrame, perimetre: str, dossier_dicos: str) -> tuple[pd.DataFrame, pd.DataFrame, dict]:
    """Table BPE au format du classeur et tableau récapitulatif (équipements par type et par zone)."""
    df_data = df_data.copy()
    df_data.insert(0, "OBJECTID", range(1, len(df_data) + 1))
    df_data, dictionnaire = _traduire(df_data, df_meta)

    if "SECT" in df_data.columns and "SECTEUR" in df_data.columns:
        df_data = df_data.drop(columns=["SECTEUR"])
    df_bpe = df_data.rename(columns=LIBELLES_COLONNES)

    colonnes = list(COLONNES_EXPORT)
    if perimetre == "roubaix":
        code_iris = df_bpe["Code IRIS"] if "Code IRIS" in df_bpe.columns else pd.Series("", index=df_bpe.index)
        for colonne, fichier in zip(COLONNES_ROUBAIX, ("secteurs_mairie.yml", os.path.join("bpe", "quartiers_historiques.yml"), "grand_quartier.yml")):
            df_bpe[colonne] = code_iris.map(_charger_yml(os.path.join(dossier_dicos, fichier))).fillna("")
    else:
        colonnes = [col for col in colonnes if col not in COLONNES_ROUBAIX]

    for col in colonnes:
        if col not in df_bpe.columns:
            df_bpe[col] = ""
    latitude = df_bpe["Latitude GPS"].fillna("").astype(str).str.strip()
    longitude = df_bpe["Longitude GPS"].fillna("").astype(str).str.strip()
    avec_coordonnees = latitude.ne("") & longitude.ne("")
    df_bpe["Latitude et longitude GPS"] = ""
    df_bpe.loc[avec_coordonnees, "Latitude et longitude GPS"] = latitude[avec_coordonnees] + ", " + longitude[avec_coordonnees]
    df_bpe = df_bpe[colonnes]

    if perimetre == "roubaix":
        libelles_iris = _charger_yml(os.path.join(dossier_dicos, "libelles_iris_roubaix.yml"))
        libelles_iris[f"{CODE_ROUBAIX}_IND"] = "NON LOCALISÉ"
        zone = df_bpe["Code IRIS"].map(libelles_iris).fillna(df_bpe["Code IRIS"]).fillna("NON LOCALISÉ")
        total = "TOTAL ROUBAIX"
    else:
        communes = _charger_yml(os.path.join(dossier_dicos, "communes_mel.yml"))
        zone = df_bpe["Code Insee"].map(communes).fillna("AUTRE / HORS MEL")
        total = "TOTAL MEL"

    recap = pd.pivot_table(
        df_bpe.assign(_zone=zone.replace("", "NON LOCALISÉ")),
        index=["Domaine d'appartenance", "Sous Domaine d'appartenance", "Type d'équipement"],
        columns="_zone",
        aggfunc="size",
        fill_value=0,
    )
    recap.columns.name = None
    recap[total] = recap.sum(axis=1)
    return df_bpe, recap.reset_index(), dictionnaire


def _ajuster_largeurs(feuille, df: pd.DataFrame) -> None:
    for index, col in enumerate(df.columns, start=1):
        longueurs = df[col].fillna("").astype(str).str.len()
        largeur = max(int(longueurs.max()) if len(longueurs) else 0, len(str(col)) // 3, 15)
        feuille.column_dimensions[get_column_letter(index)].width = min(largeur, 45)


def _formater_entetes(feuille, ligne: int = 1, couleur: str | None = BLEU) -> None:
    feuille.row_dimensions[ligne].height = 45
    for cellule in feuille[ligne]:
        cellule.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cellule.font = Font(color="FFFFFF", bold=True) if couleur else Font(bold=True)
        if couleur:
            cellule.fill = PatternFill(start_color=couleur, end_color=couleur, fill_type="solid")


def _ecrire_table(writer, df: pd.DataFrame, nom_onglet: str, nom_table: str, style: str = "TableStyleLight1", couleur: str | None = BLEU) -> None:
    df.to_excel(writer, sheet_name=nom_onglet, index=False)
    feuille = writer.sheets[nom_onglet]
    if len(df):
        table = Table(displayName=nom_table, ref=feuille.dimensions)
        table.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True)
        feuille.add_table(table)
    _ajuster_largeurs(feuille, df)
    _formater_entetes(feuille, couleur=couleur)
    feuille.freeze_panes = "A2"


def _ecrire_accueil(feuille, annee: str, perimetre: str, onglets_domaines: list[tuple[str, str]]) -> None:
    feuille.sheet_view.showGridLines = False
    for colonne, largeur in zip("ABC", (35, 120, 60)):
        feuille.column_dimensions[colonne].width = largeur

    feuille.merge_cells("A1:C1")
    titre = "Base BPE MEL - INSEE" if perimetre == "mel" else "Base BPE - INSEE"
    feuille["A1"] = f"{titre} (Mise à jour {annee})"
    feuille["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    feuille["A1"].fill = PatternFill(start_color=BLEU, end_color=BLEU, fill_type="solid")
    feuille["A1"].alignment = Alignment(horizontal="center", vertical="center")
    feuille.row_dimensions[1].height = 40

    bordure_entete = Border(bottom=Side(style="medium", color=BLEU), right=Side(style="thin", color="FFFFFF"))
    for numero, texte in enumerate(["Onglet disponible", "Description", "Lien vers le dictionnaire de données"], start=1):
        cellule = feuille.cell(row=2, column=numero, value=texte)
        cellule.font = Font(bold=True, color=BLEU_FONCE)
        cellule.fill = PatternFill(start_color=BLEU_CLAIR, end_color=BLEU_CLAIR, fill_type="solid")
        cellule.border = bordure_entete
        cellule.alignment = Alignment(horizontal="center", vertical="center")
    feuille.row_dimensions[2].height = 25

    zone = "commune de la MEL" if perimetre == "mel" else "secteur géographique"
    lignes = [
        ["BPE", "Base globale contenant l'intégralité des équipements avec leurs attributs Insee et données de géolocalisation."],
        ["Recap", f"Tableau croisé dynamique de synthèse permettant de quantifier les équipements par type et par {zone}."],
        *[[nom_onglet, DESCRIPTIONS_DOMAINES.get(code, "Détail des équipements de ce domaine.")] for code, nom_onglet in onglets_domaines],
        ["Dictionnaire Insee", "Dictionnaire des codes Insee."],
    ]
    bordure_ligne = Border(bottom=Side(style="thin", color=BLEU_CLAIR), right=Side(style="thin", color="F2F2F2"))
    for i, (onglet, description) in enumerate(lignes):
        ligne = i + 3
        fond = PatternFill(start_color="F8FAFC" if i % 2 else "FFFFFF", end_color="F8FAFC" if i % 2 else "FFFFFF", fill_type="solid")
        feuille.row_dimensions[ligne].height = 35
        for colonne, valeur in enumerate([onglet, description, LIEN_INSEE], start=1):
            cellule = feuille.cell(row=ligne, column=colonne, value=valeur)
            cellule.fill = fond
            cellule.border = bordure_ligne
            cellule.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        feuille.cell(row=ligne, column=1).font = Font(bold=True, color=BLEU_FONCE)
        feuille.cell(row=ligne, column=1).hyperlink = f"#'{onglet}'!A1"
        feuille.cell(row=ligne, column=2).font = Font(bold=True)
        feuille.cell(row=ligne, column=3).font = Font(color="0000FF", underline="single")
        feuille.cell(row=ligne, column=3).hyperlink = LIEN_INSEE

    ligne_precision = len(lignes) + 4
    feuille.merge_cells(f"A{ligne_precision}:C{ligne_precision}")
    cellule = feuille[f"A{ligne_precision}"]
    cellule.value = "Précision"
    cellule.font = Font(bold=True, color=BLEU_FONCE, size=12)
    cellule.fill = PatternFill(start_color="B8CCE4", end_color="B8CCE4", fill_type="solid")
    cellule.alignment = Alignment(horizontal="center", vertical="center")
    feuille.row_dimensions[ligne_precision].height = 25

    echelle = (
        "# Les données BPE MEL sont présentées à l'échelle communale (les 95 communes de la Métropole Européenne de Lille).\n"
        if perimetre == "mel"
        else "# Les données BPE sont présentées à l'échelle communale et infra-communale (IRIS).\n"
    )
    texte = (
        echelle
        + "Le géocodage des équipements dépend de la qualité des adresses fournies. Une marge d'erreur de positionnement est possible.\n\n"
        "# L'ensemble des variables et définitions sont présentes sur le site de l'Insee : "
        "https://www.insee.fr/fr/metadonnees/source/operation/s2216/presentation\n"
        "# Des applications dédiées sont également disponibles sur le portail SIG : https://carto.ville-roubaix.fr/\n\n"
        "Pour toutes questions, contacter le service SIG : sig@ville-roubaix.fr"
    )
    ligne_texte = ligne_precision + 1
    feuille.merge_cells(f"A{ligne_texte}:C{ligne_texte}")
    cellule = feuille[f"A{ligne_texte}"]
    cellule.value = texte
    cellule.font = Font(size=11)
    cellule.fill = PatternFill(start_color="E6EAF5", end_color="E6EAF5", fill_type="solid")
    cellule.alignment = Alignment(vertical="center", wrap_text=True)
    feuille.row_dimensions[ligne_texte].height = 130


def _ecrire_recap(writer, recap: pd.DataFrame) -> None:
    recap.to_excel(writer, sheet_name="Recap", index=False, startrow=1)
    feuille = writer.sheets["Recap"]
    derniere_ligne = feuille.max_row
    if len(recap):
        table = Table(displayName="TableRecap", ref=f"A2:{get_column_letter(feuille.max_column)}{derniere_ligne}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        feuille.add_table(table)
    _ajuster_largeurs(feuille, recap)
    _formater_entetes(feuille, ligne=2, couleur=None)

    # Ligne 1 : totaux qui suivent les filtres du tableau
    vert = Font(bold=True, color="00B050")
    feuille.row_dimensions[1].height = 30
    feuille.cell(row=1, column=3, value="Total").font = vert
    feuille.cell(row=1, column=3).alignment = Alignment(horizontal="center", vertical="center")
    for colonne in range(4, feuille.max_column + 1):
        lettre = get_column_letter(colonne)
        cellule = feuille.cell(row=1, column=colonne, value=f"=SUBTOTAL(9,{lettre}3:{lettre}{derniere_ligne})")
        cellule.font = vert
        cellule.fill = PatternFill(start_color="EBF1DE", end_color="EBF1DE", fill_type="solid")
        cellule.alignment = Alignment(horizontal="center", vertical="center")
    feuille.freeze_panes = "D3"


def generer_classeur_bpe(
    chemin_bpe: str,
    chemin_sortie: str,
    perimetre: str = "roubaix",
    chemin_varmod: str | None = None,
    dossier_dicos: str = "dictionnaires",
) -> dict:
    """Écrit le classeur BPE et renvoie un résumé (année, nombre d'équipements, par domaine…)."""
    if perimetre not in PERIMETRES:
        raise ErreurBPE(f"Périmètre inconnu : {perimetre} (attendu : {', '.join(PERIMETRES)})")
    chemin_varmod = chemin_varmod or os.path.join(dossier_dicos, FICHIER_VARMOD_DEFAUT)
    df_meta = lire_varmod(chemin_varmod)

    codes = {CODE_ROUBAIX} if perimetre == "roubaix" else set(_charger_yml(os.path.join(dossier_dicos, "communes_mel.yml")))
    df_data = extraire_perimetre(chemin_bpe, codes)
    if df_data.empty:
        raise ErreurBPE(f"Aucun équipement du périmètre {PERIMETRES[perimetre]} dans ce fichier (colonne DEPCOM).")

    annees = df_data["AN"].dropna().unique().tolist() if "AN" in df_data.columns else []
    annee = str(max(annees)) if annees else "inconnue"
    df_bpe, recap, dictionnaire = preparer_bpe(df_data, df_meta, perimetre, dossier_dicos)

    domaines = sorted(code for code in df_bpe["Code Domaine d'appartenance"].dropna().unique() if code)
    onglets_domaines = [
        (code, f"{code} - {dictionnaire.get('DOM', {}).get(code, 'Autre')}"[:31].strip()) for code in domaines
    ]

    with pd.ExcelWriter(chemin_sortie, engine="openpyxl") as writer:
        _ecrire_accueil(writer.book.create_sheet("Accueil", 0), annee, perimetre, onglets_domaines)
        _ecrire_table(writer, df_bpe, "BPE", "TableBPE")
        _ecrire_recap(writer, recap)
        for code, nom_onglet in onglets_domaines:
            df_domaine = df_bpe[df_bpe["Code Domaine d'appartenance"] == code]
            _ecrire_table(writer, df_domaine, nom_onglet, f"TableBPE_DOM_{code}")
        _ecrire_table(writer, df_meta, "Dictionnaire Insee", "TableMeta", style="TableStyleMedium2", couleur=None)

    zones_non_localisees = int((df_bpe["Code IRIS"].fillna("").str.endswith("_IND")).sum()) if perimetre == "roubaix" else 0
    return {
        "classeur": chemin_sortie,
        "annee": annee,
        "perimetre": PERIMETRES[perimetre],
        "equipements": int(len(df_bpe)),
        "communes": int(df_bpe["Code Insee"].nunique()),
        "non_localises_iris": zones_non_localisees,
        "domaines": [
            {"code": code, "libelle": dictionnaire.get("DOM", {}).get(code, "Autre"), "equipements": int((df_bpe["Code Domaine d'appartenance"] == code).sum())}
            for code, _ in onglets_domaines
        ],
        "nom_fichier": f"BPE{annee[-2:] if annee.isdigit() else ''}_{PERIMETRES[perimetre]}_equipements_officiel.xlsx",
    }
