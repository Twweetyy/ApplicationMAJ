"""Mise à jour des pages « Présentation » et « Detail calcul » selon le lot traité."""
import math
import os
import re
import unicodedata
from copy import copy

import yaml
from openpyxl.styles import Alignment, Border, Font, PatternFill

ORDRE_THEMES_CALCULS = ["POP", "LOG", "ACT RES", "COU FAM", "DIP FORMATION", "DEFM", "CAF", "CNAM", "RPLS"]
ENTETES_CALCULS = {
    "POP": "Details sur les calculs de la population",
    "LOG": "Details sur les calculs du logement",
    "ACT RES": "Details sur les calculs des activités des résidents",
    "COU FAM": "Details sur les calculs des couples-familles-ménages",
    "DIP FORMATION": "Details sur les calculs des diplômes et formations",
    "DEFM": "Details sur les calculs des demandeurs d'emploi",
    "CAF": "Details sur les calculs des bénéficiaires de la CAF",
    "CNAM": "Details sur les calculs des bénéficiaires de la CNAM",
    "RPLS": "Details sur les calculs du parc locatif social",
}
LARGEUR_COLONNE_CALCULS = 42
MOTIF_ANNEE = re.compile(r"^(.*\S)\s+(\d{4})$")


def _normaliser(texte) -> str:
    texte = unicodedata.normalize("NFKD", str(texte).casefold())
    texte = "".join(c for c in texte if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]", "", texte)


def _onglets_par_theme(dico_fichiers: dict) -> dict:
    """{'POP': ('POP 2023', '2023'), 'PETITE ENFANCE': ('PETITE ENFANCE 2022', '2022'), ...}"""
    resultat = {}
    for nom in dico_fichiers:
        correspondance = MOTIF_ANNEE.match(str(nom))
        if correspondance:
            resultat[correspondance.group(1)] = (nom, correspondance.group(2))
    return resultat


def _titre_general(onglets: dict) -> str:
    annee = onglets.get("POP", (None, None))[1] or max((a for _, a in onglets.values()), default="")
    return f"Base MAS - INSEE mise à jour {annee}".strip()


def _theme_de_libelle(libelle, themes: list[str]) -> str | None:
    """'COU FAM MÉNAGES' -> 'COU FAM' (thème le plus long dont le nom commence le libellé)."""
    cle = _normaliser(libelle)
    candidats = [theme for theme in themes if cle.startswith(_normaliser(theme))]
    return max(candidats, key=len) if candidats else None


def _charger_liens(chemin_liens: str | None) -> dict:
    try:
        with open(chemin_liens, "r", encoding="utf-8") as fichier:
            brut = yaml.safe_load(fichier) or {}
    except (OSError, TypeError):
        return {}
    return {str(theme): {str(annee): url for annee, url in (liens or {}).items()} for theme, liens in brut.items()}


def _ecrire_lien(cellule, url: str, modele_lien) -> None:
    cellule.value = url
    cellule.hyperlink = url
    if modele_lien is not None:
        cellule.font = copy(modele_lien.font)
        cellule.alignment = copy(modele_lien.alignment)
    else:
        police = copy(cellule.font)
        cellule.font = Font(name=police.name, size=police.sz, underline="single", color="0563C1")


def _maj_presentation(ws, onglets: dict, titre: str, liens: dict) -> None:
    ws["A1"] = titre
    # Une cellule déjà en lien souligné sert de modèle de style
    modele_lien = next(
        (
            ws.cell(row=r, column=5)
            for r in range(4, ws.max_row + 1)
            if ws.cell(row=r, column=5).hyperlink and ws.cell(row=r, column=5).font.u
        ),
        None,
    )
    themes = sorted(set(onglets) | set(ORDRE_THEMES_CALCULS) | {"REV DEC", "REV DISP", "PETITE ENFANCE"})
    ligne = 4
    while ws.cell(row=ligne, column=1).value not in (None, ""):
        theme = _theme_de_libelle(ws.cell(row=ligne, column=1).value, themes)
        if theme in onglets:
            annee = onglets[theme][1]
            ws.cell(row=ligne, column=2, value=int(annee))
            date = ws.cell(row=ligne, column=3).value
            if isinstance(date, str) and re.search(r"\d{4}", date):
                ws.cell(row=ligne, column=3, value=re.sub(r"\d{4}", annee, date))
            url = liens.get(theme, {}).get(annee)
            if url:
                _ecrire_lien(ws.cell(row=ligne, column=5), url, modele_lien)
        elif theme is not None:
            ws.cell(row=ligne, column=2, value="Non fourni")
        ligne += 1


def _formule_lisible(params: dict, onglets: dict) -> str:
    variables = {f"var{i}": params.get(f"var{i}") for i in range(1, 5)}
    expression = str(params.get("expression", ""))
    # Espacement des opérateurs avant de remplacer les noms (qui peuvent contenir des tirets)
    expression = re.sub(r"\s*([*/+-])\s*", r" \1 ", expression)

    def autre_onglet(m):
        theme = MOTIF_ANNEE.match(m.group(1))
        theme = theme.group(1) if theme else m.group(1)
        return f"{variables.get(m.group(2)) or m.group(2)} (feuille {theme})"

    expression = re.sub(r"""files\[\s*["']([^"']+)["']\s*\]\[\s*(var\d)\s*\]""", autre_onglet, expression)
    expression = re.sub(r"files\[\s*df\s*\]\[\s*(var\d)\s*\]", lambda m: str(variables.get(m.group(1)) or m.group(1)), expression)
    expression = expression.replace(".sum()", " (total Roubaix)")
    expression = expression.replace("( ", "(").replace(" )", ")")
    return re.sub(r"\s+", " ", expression).strip()


def _maj_detail_calcul(ws, dico_fichiers: dict, onglets: dict, chemin_calculs: str, titre: str) -> None:
    try:
        with open(chemin_calculs, "r", encoding="utf-8") as fichier:
            calculs = yaml.safe_load(fichier) or {}
    except OSError:
        return

    par_theme = {}
    for indicateur, params in calculs.items():
        correspondance = MOTIF_ANNEE.match(str(params.get("df", "")))
        theme = correspondance.group(1) if correspondance else str(params.get("df", ""))
        if theme in onglets:
            par_theme.setdefault(theme, []).append((indicateur.strip(), params))
    themes = [t for t in ORDRE_THEMES_CALCULS if t in par_theme] + sorted(t for t in par_theme if t not in ORDRE_THEMES_CALCULS)

    # Styles du gabarit (ligne 2 = en-têtes, ligne 3 = contenu) avant nettoyage
    style_entete = ws.cell(row=2, column=1)
    style_corps = ws.cell(row=3, column=1)
    styles = {
        "entete": (copy(style_entete.font), copy(style_entete.fill), copy(style_entete.border), copy(style_entete.alignment)),
        "corps": (copy(style_corps.font), copy(style_corps.fill), copy(style_corps.border), copy(style_corps.alignment)),
    }

    # Nettoyage complet (valeurs, styles, hauteurs) sous le titre
    derniere_colonne = max(ws.max_column, len(themes))
    for ligne in range(2, ws.max_row + 1):
        ws.row_dimensions[ligne].height = None
        for colonne in range(1, derniere_colonne + 1):
            cellule = ws.cell(row=ligne, column=colonne)
            cellule.value = None
            cellule.font, cellule.border, cellule.fill = Font(), Border(), PatternFill()
            cellule.alignment = Alignment()

    for merged in list(ws.merged_cells.ranges):
        if merged.min_row == 1:
            ws.unmerge_cells(str(merged))
    ws["A1"] = titre
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(themes), 1))

    def appliquer(cellule, cle):
        police, fond, bordure, alignement = styles[cle]
        cellule.font, cellule.fill, cellule.border = copy(police), copy(fond), copy(bordure)
        cellule.alignment = Alignment(horizontal=alignement.horizontal or "center", vertical="center", wrap_text=True)

    for colonne, theme in enumerate(themes, start=1):
        entete = ws.cell(row=2, column=colonne, value=ENTETES_CALCULS.get(theme, f"Details sur les calculs {theme}"))
        appliquer(entete, "entete")
        ws.column_dimensions[entete.column_letter].width = LARGEUR_COLONNE_CALCULS
        colonnes_finales = set(dico_fichiers[onglets[theme][0]].columns)
        for ligne, (indicateur, params) in enumerate(par_theme[theme], start=3):
            texte = f"{indicateur} -> {_formule_lisible(params, onglets)}"
            calcule = indicateur in colonnes_finales
            if not calcule:
                texte += "\n(non calculé : variable absente de ce millésime)"
            cellule = ws.cell(row=ligne, column=colonne, value=texte)
            appliquer(cellule, "corps")
            if not calcule:
                police = copy(cellule.font)
                cellule.font = Font(name=police.name, size=police.sz, italic=True, color="C00000")

    # Hauteurs de lignes estimées (Excel ne les recalcule pas pour le texte renvoyé à la ligne)
    for ligne in range(2, ws.max_row + 1):
        hauteur = 15.0
        for colonne in range(1, len(themes) + 1):
            cellule = ws.cell(row=ligne, column=colonne)
            if not cellule.value:
                continue
            taille = cellule.font.sz or 11
            caracteres_par_ligne = max(1, int(LARGEUR_COLONNE_CALCULS * 1.15 * 11 / taille))
            lignes = sum(math.ceil(max(len(morceau), 1) / caracteres_par_ligne) for morceau in str(cellule.value).split("\n"))
            hauteur = max(hauteur, lignes * taille * 1.35 + 6)
        ws.row_dimensions[ligne].height = min(hauteur, 409)


def mettre_a_jour_pages_info(
    workbook,
    dico_fichiers: dict,
    chemin_calculs: str | None,
    chemin_liens: str | None = None,
) -> None:
    """Adapte les deux premières pages du classeur au lot réellement traité."""
    onglets = _onglets_par_theme(dico_fichiers)
    titre = _titre_general(onglets)
    if chemin_liens is None and chemin_calculs:
        chemin_liens = os.path.join(os.path.dirname(chemin_calculs), "liens_presentation.yml")
    if "Présentation" in workbook.sheetnames:
        _maj_presentation(workbook["Présentation"], onglets, titre, _charger_liens(chemin_liens))
    if "Detail calcul" in workbook.sheetnames and chemin_calculs:
        _maj_detail_calcul(workbook["Detail calcul"], dico_fichiers, onglets, chemin_calculs, titre)
