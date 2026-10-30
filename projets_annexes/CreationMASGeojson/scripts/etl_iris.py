"""
ETL IRIS Roubaix : TAB_BDDMAS_Tout_perimetre_<ANNEE>.xlsx  ->  schéma stat_insee_iris2 (champ_insee_N)
Sorties :
  data/traite/stat_insee_iris_<ANNEE>.xlsx   (même format que stat_insee_iris2 : 2 lignes d'en-tête)
  data/traite/roubaix_iris_<ANNEE>.geojson   (38 polygones IRIS + indicateurs)

Usage :  python scripts/etl_iris.py            (année par défaut ci-dessous)
         python scripts/etl_iris.py 2024       (millésime suivant, mêmes fichiers de référence)
Dépendances : pandas, openpyxl, pyyaml, xlrd (lecture du .xls modèle)
"""
import json
import logging
import math
import sys
import unicodedata
import warnings
from pathlib import Path

import pandas as pd
import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------- config
ANNEE = int(sys.argv[1]) if len(sys.argv) > 1 else 2023

RACINE = Path(__file__).resolve().parents[1]
BRUT, REF, TRAITE = RACINE / "data/brut", RACINE / "data/ref", RACINE / "data/traite"
FICHIER_BDDMAS = BRUT / f"TAB_BDDMAS_Tout_perimetre_{ANNEE}.xlsx"
FICHIER_GEOJSON = BRUT / "Iris_2000.geojson"
FICHIER_MODELE = BRUT / "stat_insee_iris2.xls"          # modèle de colonnes (+ historique au 1er lancement)
# Historique : la sortie du millésime précédent si elle existe (elle contient déjà ce millésime),
# sinon le modèle stat_insee_iris2.xls.
FICHIER_HISTO_PREC = TRAITE / f"stat_insee_iris_{ANNEE - 1}.xlsx"
FICHIER_MAPPING = REF / "mapping_champs_insee.yml"
# Champs calculés recalculés aussi sur l'historique (définition corrigée, série homogène)
RECALCUL_HISTO = ["champ_insee_91"]
# Colonnes techniques ArcGIS, jamais exportées dans le GeoJSON
COLS_TECHNIQUES = ["objectid", "created_date", "last_edited_date", "date_stat", "date_statistique",
                   "st_area(shape)", "st_length(shape)"]
CODE_COMMUNE = "59512"
VALEURS_SECRET = ["s", "ns", "nd", "-", ""]

logging.basicConfig(level=logging.INFO, format="%(levelname)s | %(message)s")
log = logging.getLogger("etl_iris")
warnings.simplefilter("ignore", pd.errors.PerformanceWarning)


def norm(txt) -> str:
    """Normalise un libellé : sans accents, apostrophes droites, minuscules, espaces simples."""
    txt = str(txt).replace("’", "'").replace(" ", " ").replace("²", "2")
    txt = "".join(c for c in unicodedata.normalize("NFKD", txt) if not unicodedata.combining(c))
    return " ".join(txt.lower().split())


# --------------------------------------------------------------------------- 1. extraction
def lire_geojson():
    geo = json.loads(FICHIER_GEOJSON.read_text(encoding="utf-8"))
    ref = pd.DataFrame([f["properties"] for f in geo["features"]])
    ref = ref.rename(columns={"numeros": "id_rbx"})[["id_rbx", "libelle", "code_iris", "nom_iris"]]
    ref["id_rbx"] = ref["id_rbx"].astype(str).str.zfill(4)
    ref["code_iris"] = ref["code_iris"].astype(str)
    assert (ref["code_iris"] == CODE_COMMUNE + ref["id_rbx"]).all(), "code_iris incohérent avec numeros"
    ref["cle"] = ref["nom_iris"].map(norm)
    log.info("GeoJSON : %d IRIS", len(ref))
    return geo, ref


def lire_onglets():
    """Retourne {ALIAS_ONGLET: DataFrame indexé par libellé IRIS normalisé}. ALIAS = nom d'onglet sans l'année."""
    xls = pd.ExcelFile(FICHIER_BDDMAS)
    onglets, noms, mdq = {}, {}, None
    for nom in xls.sheet_names:
        if nom in ("Présentation", "Detail calcul"):
            continue
        alias = " ".join(m for m in nom.split() if not m.isdigit())
        df = pd.read_excel(xls, nom, header=1)
        df = df[df["Libellé de l'Iris"].notna()]
        df.index = df["Libellé de l'Iris"].map(norm)
        df.columns = [norm(c) for c in df.columns]
        onglets[alias], noms[alias] = df, nom
        if mdq is None and "mairies de quartiers" in df.columns:
            mdq = df["mairies de quartiers"]
        log.info("Onglet %-22s -> alias %-14s %d lignes, %d colonnes", nom, alias, len(df), df.shape[1])
    return onglets, noms, mdq


# --------------------------------------------------------------------------- 2. transformation
def evaluer(expr, df):
    champs = {c: pd.to_numeric(df[c], errors="coerce") for c in df.columns if str(c).startswith("champ_insee_")}
    return pd.eval(expr, local_dict=champs)


def construire_table(ref, onglets, noms, mdq, mapping):
    out = ref.set_index("cle").copy()
    out["perim"] = out["libelle"]
    out["perim_mdq"] = mdq.reindex(out.index).values
    rapport = []

    # a) sources directes + YAML
    for champ, regle in mapping.items():
        if "src" in regle:
            alias, col = regle["src"]
            df = onglets.get(alias)
            if df is None or norm(col) not in df.columns:
                raise KeyError(f"{champ}: colonne '{col}' introuvable dans l'onglet '{alias}'")
            s = df[norm(col)].replace(VALEURS_SECRET, pd.NA)
            out[champ] = pd.to_numeric(s, errors="coerce").reindex(out.index)
            rapport.append((champ, f"{noms[alias]} | {col}", "OK"))   # nom réel d'onglet = millésime de la source
        elif "yml" in regle:
            d = yaml.safe_load((REF / regle["yml"]).read_text(encoding="utf-8"))
            out[champ] = out["libelle"].map({str(k).strip(): v for k, v in d.items()})
            rapport.append((champ, f"YAML {regle['yml']}", "OK"))
        elif "na" in regle:
            out[champ] = float("nan")
            rapport.append((champ, regle["na"], "NON DISPONIBLE"))
        elif "supprimer" in regle:
            rapport.append((champ, regle["supprimer"], "SUPPRIMÉ"))

    # b) champs calculés (ordre du YAML)
    for champ, regle in mapping.items():
        if "expr" in regle:
            out[champ] = evaluer(regle["expr"], out)
            statut = "ESTIMATION" if "note" in regle else "OK"
            source = f"Calcul : {regle['expr']}" + (f" ({regle['note']})" if "note" in regle else "")
            rapport.append((champ, source, statut))

    out["annee"] = ANNEE
    rapport.sort(key=lambda r: int(r[0].split("_")[-1]))
    return out.reset_index(drop=True), pd.DataFrame(rapport, columns=["champ", "source", "statut"])


# --------------------------------------------------------------------------- 3. mise au format du modèle
def lire_modele(mapping):
    """Libellés et alias toujours lus dans le modèle (schéma complet, même si une sortie a perdu des colonnes vides).
    Le mapping peut remplacer un libellé (clé libelle) et ajouter des champs absents du modèle (mis en fin de table).
    Historique lu dans la sortie du millésime précédent si elle existe, sinon dans le modèle."""
    entete = pd.read_excel(FICHIER_MODELE, header=None, nrows=2)
    libelles, alias = entete.iloc[0].tolist(), entete.iloc[1].tolist()
    nouveaux = [c for c in mapping if c not in alias]
    alias += nouveaux
    libelles += nouveaux
    libelles = [mapping.get(a, {}).get("libelle", l) for l, a in zip(libelles, alias)]
    garder = [i for i, a in enumerate(alias) if "supprimer" not in mapping.get(a, {})]
    libelles, alias = [libelles[i] for i in garder], [alias[i] for i in garder]
    if FICHIER_HISTO_PREC.exists():
        source, onglet = FICHIER_HISTO_PREC, "historique_complet"
    else:
        source, onglet = FICHIER_MODELE, 0
    brut = pd.read_excel(source, sheet_name=onglet, header=None)
    histo = brut.iloc[2:].copy()
    histo.columns = brut.iloc[1].tolist()
    histo = histo.reindex(columns=alias)
    histo["id_rbx"] = histo["id_rbx"].astype(str).str.zfill(4)
    histo["code_iris"] = histo["code_iris"].astype(str)
    histo["annee"] = pd.to_numeric(histo["annee"]).astype(int)
    for champ in RECALCUL_HISTO:
        histo[champ] = evaluer(mapping[champ]["expr"], histo)
    log.info("Historique lu dans %s : %d lignes (%s)", source.name, len(histo),
             ", ".join(sorted(histo["annee"].astype(str).unique())))
    return libelles, alias, histo


def mettre_au_format(table, alias, histo):
    t = table.copy()
    t["objectid"] = range(int(pd.to_numeric(histo["objectid"]).max()) + 1,
                          int(pd.to_numeric(histo["objectid"]).max()) + 1 + len(t))
    geom = histo.sort_values("annee").groupby("id_rbx")[["st_area(shape)", "st_length(shape)"]].last()
    t = t.join(geom, on="id_rbx")
    for c in alias:
        if c not in t.columns:
            t[c] = None                       # created_date, last_edited_date, date_statistique...
    return t[alias].sort_values("id_rbx", ascending=False).reset_index(drop=True)


# --------------------------------------------------------------------------- 4. contrôles
def controles(t, histo, rapport):
    lignes = []
    lignes.append(("Nombre d'IRIS", len(t), "OK" if len(t) == 38 else "ERREUR"))
    lignes.append(("id_rbx manquants", int(t["id_rbx"].isna().sum()), "OK" if t["id_rbx"].notna().all() else "ERREUR"))
    lignes.append(("perim_mdq manquants", int(t["perim_mdq"].isna().sum()), "OK" if t["perim_mdq"].notna().all() else "ERREUR"))
    lignes.append((f"Population totale Roubaix {ANNEE}", int(t["champ_insee_2"].sum()), "info"))
    # Comparaison avec la dernière année de l'historique : écart médian par champ
    der = histo[histo["annee"].astype(str) == str(pd.to_numeric(histo["annee"], errors="coerce").max())]
    der = der.set_index("id_rbx")
    cur = t.set_index("id_rbx")
    for champ in rapport.loc[rapport.statut == "OK", "champ"]:
        a = pd.to_numeric(der[champ], errors="coerce")
        b = pd.to_numeric(cur[champ], errors="coerce")
        ecart = ((b - a).abs() / a.abs().replace(0, pd.NA)).dropna()
        if len(ecart):
            med = float(ecart.median())
            if med > 0.35:
                lignes.append((f"Écart médian {champ} vs {int(der['annee'].iloc[0])}", round(med * 100, 1), "À VÉRIFIER"))
    return pd.DataFrame(lignes, columns=["contrôle", "valeur", "statut"])


# --------------------------------------------------------------------------- 5. chargement
ENTETE_FILL = PatternFill("solid", fgColor="1F4E78")
ALIAS_FILL = PatternFill("solid", fgColor="D9E1F2")
TEXTE = {"id_rbx", "code_iris"}


def colonnes_non_vides(df):
    """Colonnes qui ont au moins une valeur (les colonnes totalement vides ne sont pas exportées)."""
    return [c for c in df.columns if df[c].replace("", pd.NA).notna().any()]


def ecrire_feuille(ws, libelles, alias, df):
    garder = colonnes_non_vides(df)
    lib = dict(zip(alias, libelles))
    alias = [a for a in alias if a in garder]
    libelles = [lib[a] for a in alias]
    df = df[alias]
    ws.append(libelles)
    ws.append(alias)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), ENTETE_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for c in ws[2]:
        c.font, c.fill = Font(bold=True, italic=True), ALIAS_FILL
    for row in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and math.isnan(v)) or v is pd.NA else v for v in row])
    for j, a in enumerate(alias, start=1):
        if a in TEXTE:
            for r in range(3, ws.max_row + 1):
                ws.cell(r, j).number_format = "@"
                ws.cell(r, j).value = None if ws.cell(r, j).value is None else str(ws.cell(r, j).value)
        ws.column_dimensions[get_column_letter(j)].width = 14 if j > 7 else 20
    ws.row_dimensions[1].height = 60
    ws.freeze_panes = "H3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(alias))}{ws.max_row}"


def ecrire_excel(libelles, alias, t, histo, rapport, ctrl):
    chemin = TRAITE / f"stat_insee_iris_{ANNEE}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = f"stat_insee_{ANNEE}"
    ecrire_feuille(ws, libelles, alias, t)

    h = histo[histo["annee"].astype(str) != str(ANNEE)]
    complet = pd.concat([t, h[alias]], ignore_index=True)
    ecrire_feuille(wb.create_sheet("historique_complet"), libelles, alias, complet)

    lib = dict(zip(alias, libelles))
    rap = rapport.assign(libelle=rapport["champ"].map(lib))[["champ", "libelle", "source", "statut"]]
    for nom, df in (("mapping_sources", rap), ("controles", ctrl)):
        w = wb.create_sheet(nom)
        w.append(list(df.columns))
        for c in w[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), ENTETE_FILL
        for row in df.itertuples(index=False):
            w.append(list(row))
        for j in range(1, df.shape[1] + 1):
            w.column_dimensions[get_column_letter(j)].width = [18, 60, 80, 16][min(j - 1, 3)]
        w.auto_filter.ref = w.dimensions
    wb.save(chemin)
    log.info("Excel écrit : %s", chemin)
    return chemin


def ecrire_geojson(geo, t, noms=None, suffixe=""):
    """noms : {alias: libellé} pour nommer les champs par leur libellé lisible au lieu de champ_insee_N."""
    chemin = TRAITE / f"roubaix_iris_{ANNEE}{suffixe}.geojson"
    noms = noms or {}
    props = t.drop(columns=COLS_TECHNIQUES, errors="ignore")
    props = props[colonnes_non_vides(props)].set_index("id_rbx")
    feats, manquants = [], []
    for f in geo["features"]:
        idr = str(f["properties"]["numeros"]).zfill(4)
        if idr not in props.index:
            manquants.append(idr)
            continue
        p = {"id_rbx": idr}
        for k, v in props.loc[idr].items():
            if v is pd.NA or (isinstance(v, float) and math.isnan(v)):
                v = None
            elif hasattr(v, "item"):
                v = v.item()
            p[noms.get(k, k)] = v
        feats.append({"type": "Feature", "properties": p, "geometry": f["geometry"]})
    if manquants:
        log.warning("IRIS du GeoJSON sans données : %s", manquants)
    sortie = {"type": "FeatureCollection", "name": f"roubaix_iris_{ANNEE}{suffixe}", "crs": geo.get("crs"), "features": feats}
    chemin.write_text(json.dumps(sortie, ensure_ascii=False), encoding="utf-8")
    log.info("GeoJSON écrit : %s (%d entités)", chemin, len(feats))
    return chemin


# --------------------------------------------------------------------------- main
def main():
    TRAITE.mkdir(parents=True, exist_ok=True)
    geo, ref = lire_geojson()
    onglets, noms, mdq = lire_onglets()
    sans_donnees = set(ref["cle"]) - set(onglets["POP"].index)
    if sans_donnees:
        raise ValueError(f"IRIS du GeoJSON absents de l'Excel : {sans_donnees}")
    mapping = yaml.safe_load(FICHIER_MAPPING.read_text(encoding="utf-8"))
    table, rapport = construire_table(ref, onglets, noms, mdq, mapping)
    libelles, alias, histo = lire_modele(mapping)
    t = mettre_au_format(table, alias, histo)
    ctrl = controles(t, histo, rapport)
    log.info("Champs non disponibles en %d : %s", ANNEE,
             ", ".join(rapport.loc[rapport.statut == "NON DISPONIBLE", "champ"]))
    print(ctrl.to_string(index=False))
    ecrire_excel(libelles, alias, t, histo, rapport, ctrl)
    ecrire_geojson(geo, t)
    # Variante avec les libellés de stat_insee_iris (1re ligne d'en-tête) comme noms de champs
    ecrire_geojson(geo, t, noms={a: str(l).strip() for a, l in zip(alias, libelles)}, suffixe="_libelles")


if __name__ == "__main__":
    main()
