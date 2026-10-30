"""
Portrait des Comités de Quartier (CQ) de Roubaix à partir de TAB_BDDMAS_Tout_perimetre_<ANNEE>.xlsx
Sorties :
  data/traite/portrait_cq_<ANNEE>.xlsx     (onglets Démographie / Socio-économie / Logement / Équipements + méthode + rattachement)
  data/traite/cq_roubaix_<ANNEE>.geojson   (polygones des CQ = fusion des IRIS, avec les indicateurs)

Usage :  python scripts/portrait_cq.py           (2023)
         python scripts/portrait_cq.py 2024
Équipements : fichier BPE de Roubaix (Insee) déposé dans data/brut sous le nom BPE<AA>_Roubaix*.xlsx (le plus récent est pris).
Dépendances : pandas, openpyxl, pyyaml, shapely
"""
import json
import logging
import math
import sys
import unicodedata
from pathlib import Path

import pandas as pd
import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from shapely.geometry import mapping, shape
from shapely.ops import unary_union

ANNEE = int(sys.argv[1]) if len(sys.argv) > 1 else 2023

RACINE = Path(__file__).resolve().parents[1]
REF, TRAITE = RACINE / "data/ref", RACINE / "data/traite"
# Projet voisin CreationMASGeojson : il fournit le TAB, les contours IRIS et stat_insee_iris (sortie de etl_iris.py).
# Un fichier présent dans ce projet (data/brut ou data/traite) est pris en priorité.
PROJET_MAS = RACINE.parent / "CreationMASGeojson"


def chercher(nom, *sous_dossiers):
    essais = [base / "data" / sd for base in (RACINE, PROJET_MAS) for sd in sous_dossiers]
    for dossier in essais:
        if (dossier / nom).exists():
            return dossier / nom
    raise FileNotFoundError(f"{nom} introuvable dans : " + " ; ".join(map(str, essais)))


def chercher_bpe():
    """Fichier BPE le plus récent (BPE<AA>_Roubaix*.xlsx) dans data/brut."""
    trouves = sorted(f for base in (RACINE, PROJET_MAS) for f in (base / "data/brut").glob("BPE*_Roubaix*.xlsx"))
    if not trouves:
        raise FileNotFoundError("Fichier BPE introuvable : déposer BPE<AA>_Roubaix_equipements_officiel.xlsx dans data/brut")
    return trouves[-1]


FICHIER_CQ = REF / "cq_iris.yml"
FICHIER_COMMUNE = REF / "roubaix_commune.yml"            # indicateurs communaux plus récents (facultatif)
HORS_CQ, ROUBAIX = "Hors CQ", "Roubaix"
# Les IRIS hors CQ restent dans le total Roubaix mais n'ont pas de ligne ni de contour propres
AFFICHER_HORS_CQ = False
VALEURS_SECRET = ["s", "ns", "nd", "-", ""]

logging.basicConfig(level=logging.INFO, format="%(levelname)s | %(message)s")
log = logging.getLogger("portrait_cq")


def norm(txt) -> str:
    txt = str(txt).replace("’", "'").replace("\xa0", " ").replace("²", "2")
    txt = "".join(c for c in unicodedata.normalize("NFKD", txt) if not unicodedata.combining(c))
    return " ".join(txt.lower().split())


# --------------------------------------------------------------------------- indicateurs
# Chaque indicateur est calculé sur les effectifs ADDITIONNÉS des IRIS du CQ (jamais une moyenne de taux).
#   ("somme", A)            : somme de la colonne A
#   ("part", [A, B...], D)  : (A + B + ...) / D * 100
#   ("ratio", A, D, k)      : A / D * k
#   ("poids", A)            : somme de A dans la zone / somme de A dans Roubaix x 100
#   ("approche", A, P)      : moyenne de A pondérée par P sur les IRIS renseignés (revenus Filosofi : non additifs)
# A = "ONGLET|Nom de colonne" (onglet sans l'année).
P, AR, CF, DF, LG, RD, CA, CN, DE, RP = ("POP", "ACT RES", "COU FAM", "DIP FORMATION", "LOG",
                                        "REV DISP", "CAF", "CNAM", "DEFM", "RPLS")
BE = "BPE"                     # Base permanente des équipements : nombre d'équipements par groupe ci-dessous


def c(onglet, col):
    return f"{onglet}|{col}"


POPU = c(P, "Population")
AGES = ["0-2 ans", "3-5 ans", "6-10 ans", "11-17 ans", "18-24 ans", "25-39 ans", "40-54 ans",
        "55-64 ans", "65-79 ans", "80 ans ou plus"]
EPOQUES = ["avt 1919", "1919 à 1945", "1946 à 1970", "1971 à 1990", "1991 à 2005", "2006 à 2019"]
RPRINC = c(LG, "Résidences principales")
FAM = c(CF, "Familles")
# Groupes d'équipements BPE : codes « Type d'équipement » Insee ou préfixes (« A1 » = tous les A1xx).
# Dictionnaire des codes : onglet « Dictionnaire Insee » du fichier BPE. Un groupe ajouté ici doit aussi
# être ajouté dans INDICATEURS (et CHIFFRES_CLES pour les pages récap).
BPE_GROUPES = {
    "Écoles maternelles et élémentaires": ["C107", "C108", "C109"],
    "Collèges": ["C201"],
    "Lycées": ["C301", "C302"],
    "Crèches et micro-crèches": ["D502", "D509"],
    "Équipements culturels": ["F3"],                       # cinéma, conservatoire, bibliothèque, spectacle, patrimoine...
    "Équipements sportifs": ["F1"],
    "Gymnases et salles multisports": ["F121"],
    "Plateaux et terrains extérieurs": ["F111", "F113"],
    "Centres sociaux": ["D506"],
    "Services publics": ["A1"],                            # mairie, France Services, France Travail, police, justice, impôts...
    "Bureaux et relais de poste": ["A206", "A207", "A208"],
    "Commerces alimentaires de proximité": ["B201", "B202", "B204", "B207", "B208", "B209", "B210"],
    "Supermarchés et hypermarchés": ["B104", "B105"],
}

INDICATEURS = [
    # ------------------------------------------------------------------ Démographie
    ("Démographie", "Population", ("somme", POPU)),
    ("Démographie", "Population 2016", ("somme", "STAT|champ_insee_7")),
    ("Démographie", "Évolution de la population depuis 2016 (%)", ("evol", POPU, "STAT|champ_insee_7")),
    ("Démographie", "Population 2006", ("somme", "STAT|champ_insee_3")),
    ("Démographie", "Évolution de la population depuis 2006 (%)", ("evol", POPU, "STAT|champ_insee_3")),
    ("Démographie", "Superficie (km²)", ("somme", c(P, "Superficie en km²"))),
    ("Démographie", "Densité (hab/km²)", ("ratio", POPU, c(P, "Superficie en km²"), 1)),
    ("Démographie", "Part de la population de Roubaix (%)", ("poids", POPU)),
    *[("Démographie", f"Part des {a} (%)", ("part", [c(P, f"Pop {a}")], POPU)) for a in AGES],
    ("Démographie", "Part des moins de 20 ans (%)", ("part", [c(P, "Pop 0-19 ans")], POPU)),
    ("Démographie", "Part des 65 ans ou plus (%)", ("part", [c(P, "Pop 65 ans ou plus")], POPU)),
    ("Démographie", "Indice de vieillissement (65 ans ou plus / moins de 20 ans x 100)",
     ("ratio", c(P, "Pop 65 ans ou plus"), c(P, "Pop 0-19 ans"), 100)),
    ("Démographie", "Part des femmes (%)", ("part", [c(P, "Pop Femmes")], POPU)),
    ("Démographie", "Part des hommes (%)", ("part", [c(P, "Pop Hommes")], POPU)),
    ("Démographie", "Ménages", ("somme", c(CF, "Ménages"))),
    ("Démographie", "Taille moyenne des ménages", ("ratio", c(LG, "Pop Ménages"), c(LG, "Ménages"), 1)),
    ("Démographie", "Part des ménages d'une personne (%)", ("part", [c(CF, "Ménages 1 personne")], c(CF, "Ménages"))),
    ("Démographie", "Familles", ("somme", FAM)),
    ("Démographie", "Part des familles monoparentales (%)", ("part", [c(CF, "Fam Monoparentales")], FAM)),
    ("Démographie", "Part des couples avec enfant(s) (%)", ("part", [c(CF, "Fam Couple avec enfant(s)")], FAM)),
    ("Démographie", "Part des couples sans enfant (%)", ("part", [c(CF, "Fam Couple sans enfant")], FAM)),
    ("Démographie", "Part des familles nombreuses (3 enfants ou plus de moins de 25 ans) (%)",
     ("part", [c(CF, "Fam 3 enfants moins 25 ans"), c(CF, "Fam 4 enfants ou plus moins 25 ans")], FAM)),
    # ------------------------------------------------------------------ Socio-économie
    ("Socio-économie", "Taux de pauvreté au seuil de 60 % (%) (approché)",
     ("approche", c(RD, "Taux de pauvreté au seuil de 60 % (%)"), c(LG, "Pop Ménages"))),
    ("Socio-économie", "Revenu disponible médian par UC (€) (approché)",
     ("approche", c(RD, "Médiane (€)"), c(LG, "Pop Ménages"))),
    ("Socio-économie", "Part des prestations sociales dans le revenu disponible (%) (approché)",
     ("approche", c(RD, "Part de l'ensemble des prestations sociales (%)"), c(LG, "Pop Ménages"))),
    ("Socio-économie", "Population 15-64 ans", ("somme", c(AR, "Pop 15-64 ans"))),
    ("Socio-économie", "Actifs 15-64 ans", ("somme", c(AR, "Actifs 15-64 ans"))),
    ("Socio-économie", "Inactifs 15-64 ans", ("somme", c(AR, "Inactifs 15-64 ans"))),
    ("Socio-économie", "Part des inactifs parmi les 15-64 ans (%)", ("part", [c(AR, "Inactifs 15-64 ans")], c(AR, "Pop 15-64 ans"))),
    ("Socio-économie", "Taux d'activité des 15-64 ans (%)", ("part", [c(AR, "Actifs 15-64 ans")], c(AR, "Pop 15-64 ans"))),
    ("Socio-économie", "Taux d'emploi des 15-64 ans (%)", ("part", [c(AR, "Actifs occupés 15-64 ans")], c(AR, "Pop 15-64 ans"))),
    ("Socio-économie", "Taux d'emploi des femmes 15-64 ans (%)",
     ("part", [c(AR, "Actifs occupés 15-64 ans Femmes")], c(AR, "Pop 15-64 ans Femmes"))),
    ("Socio-économie", "Chômeurs 15-64 ans", ("somme", c(AR, "Chômeurs 15-64 ans"))),
    ("Socio-économie", "Taux de chômage des 15-64 ans (%)", ("part", [c(AR, "Chômeurs 15-64 ans")], c(AR, "Actifs 15-64 ans"))),
    ("Socio-économie", "Taux de chômage des 15-24 ans (%)", ("part", [c(AR, "Chômeurs 15-24 ans")], c(AR, "Actifs 15-24 ans"))),
    ("Socio-économie", "Part des actifs occupés en CDI ou fonction publique (%)",
     ("part", [c(AR, "Salariés 15 ans ou plus Fonct publ, CDI")], c(AR, "Actifs occupés 15 ans ou plus"))),
    ("Socio-économie", "Part des actifs occupés à temps partiel (%)",
     ("part", [c(AR, "Actifs occ 15 ans ou plus TP")], c(AR, "Actifs occupés 15 ans ou plus"))),
    *[("Socio-économie", f"Part des {lib} parmi les 15 ans ou plus (%)", ("part", [c(P, f"Pop 15 ans ou plus {col}")], c(P, "Pop 15 ans ou plus")))
      for lib, col in [("cadres", "Cadres, Prof. intel. sup."), ("professions intermédiaires", "Prof. intermédiaires"),
                       ("employés", "Employés"), ("ouvriers", "Ouvriers"), ("retraités", "Retraités"),
                       ("autres sans activité professionnelle", "Autres")]],
    ("Socio-économie", "Part des non-scolarisés de 15 ans ou plus sans diplôme ou CEP (%)",
     ("part", [c(DF, "Pop 15 ans ou plus non scol. Sans diplôme ou CEP")], c(DF, "Pop 15 ans ou plus non scolarisée"))),
    ("Socio-économie", "Part des non-scolarisés de 15 ans ou plus diplômés du supérieur (%)",
     ("part", [c(DF, "Pop 15 ans ou plus non scol. Enseignement sup de niveau bac + 2"),
               c(DF, "Pop 15 ans ou plus non scol. Enseignement sup de niveau bac + 3 ou 4"),
               c(DF, "Pop 15 ans ou plus non scol. Enseignement sup de niveau bac + 5 ou plus")],
      c(DF, "Pop 15 ans ou plus non scolarisée"))),
    ("Socio-économie", "Taux de scolarisation des 18-24 ans (%)",
     ("part", [c(DF, "Pop scolarisée 18-24 ans")], c(DF, "Pop 18-24 ans"))),
    ("Socio-économie", "Allocataires CAF", ("somme", c(CA, "Nombre total d'allocataires"))),
    ("Socio-économie", "Part de la population couverte par la CAF (%)",
     ("part", [c(CA, "Personnes couvertes")], c(CA, "Population (CAF)"))),
    ("Socio-économie", "Allocataires au RSA socle", ("somme", c(CA, "Allocataires percevant le RSA socle"))),
    ("Socio-économie", "Part des allocataires au RSA socle (%)",
     ("part", [c(CA, "Allocataires percevant le RSA socle")], c(CA, "Nombre total d'allocataires"))),
    ("Socio-économie", "Part des allocataires à l'AAH (%)",
     ("part", [c(CA, "Allocataires percevant l’Allocation Adulte Handicapé")], c(CA, "Nombre total d'allocataires"))),
    ("Socio-économie", "Part des allocataires monoparents (%)",
     ("part", [c(CA, "Allocataires mono-parent")], c(CA, "Nombre total d'allocataires"))),
    ("Socio-économie", "Part de la population bénéficiaire de la C2S non participative (%)",
     ("part", [c(CN, "Population couverte bénéficiaire de la C2S non participative")], c(CN, "Population couverte"))),
    ("Socio-économie", "Demandeurs d'emploi cat. ABC", ("somme", c(DE, "DEFM catégories ABC"))),
    # DEFM (France Travail, fin d'année) rapportés à la population 15-64 ans du recensement : millésimes différents
    ("Socio-économie", "Part des demandeurs d'emploi ABC parmi les 15-64 ans (%)",
     ("part", [c(DE, "DEFM catégories ABC")], c(AR, "Pop 15-64 ans"))),
    ("Socio-économie", "Part des demandeurs ABC inscrits depuis 2 ans ou plus (%)",
     ("part", [c(DE, "DEFM catégories ABC d'ancienneté au chômage d'au moins 2 ans")], c(DE, "DEFM catégories ABC"))),
    ("Socio-économie", "Part des demandeurs ABC de moins de 26 ans (%)",
     ("part", [c(DE, "DEFM catégories ABC moins de 26 ans")], c(DE, "DEFM catégories ABC"))),
    ("Socio-économie", "Part des demandeurs ABC bénéficiaires du RSA (%)",
     ("part", [c(DE, "DEFM catégories ABC bénéficiaires du RSA")], c(DE, "DEFM catégories ABC"))),
    # ------------------------------------------------------------------ Logement
    ("Logement", "Logements", ("somme", c(LG, "Logements"))),
    ("Logement", "Résidences principales", ("somme", RPRINC)),
    ("Logement", "Part des logements vacants (%)", ("part", [c(LG, "Logements vacants")], c(LG, "Logements"))),
    ("Logement", "Maisons", ("somme", c(LG, "Maisons"))),
    ("Logement", "Appartements", ("somme", c(LG, "Appartements"))),
    ("Logement", "Part des maisons (%)", ("part", [c(LG, "Maisons")], c(LG, "Logements"))),
    ("Logement", "Part des appartements (%)", ("part", [c(LG, "Appartements")], c(LG, "Logements"))),
    ("Logement", "Part des propriétaires occupants (%)", ("part", [c(LG, "Rés princ occupées Propriétaires")], RPRINC)),
    ("Logement", "Part des locataires (%)", ("part", [c(LG, "Rés princ occupées Locataires")], RPRINC)),
    ("Logement", "Part des locataires HLM (%)", ("part", [c(LG, "Rés princ HLM louée vide")], RPRINC)),
    ("Logement", "Part des locataires du parc privé (%)",
     ("part_diff", c(LG, "Rés princ occupées Locataires"), c(LG, "Rés princ HLM louée vide"), RPRINC)),
    ("Logement", "Logements sociaux (RPLS)", ("somme", c(RP, "Nombre de logements du Parc Locatif Social"))),
    ("Logement", "Part des RP de 1 ou 2 pièces (%)", ("part", [c(LG, "Rés princ 1 pièce"), c(LG, "Rés princ 2 pièces")], RPRINC)),
    ("Logement", "Part des RP de 3 ou 4 pièces (%)", ("part", [c(LG, "Rés princ 3 pièces"), c(LG, "Rés princ 4 pièces")], RPRINC)),
    ("Logement", "Part des RP de 5 pièces ou plus (%)", ("part", [c(LG, "Rés princ 5 pièces ou plus")], RPRINC)),
    ("Logement", "Part des RP de moins de 40 m² (%)",
     ("part", [c(LG, "Rés princ de moins de 30 m2"), c(LG, "Rés princ de 30 à moins de 40 m2")], RPRINC)),
    *[("Logement", f"Part des RP construites {e.replace('avt', 'avant')} (%)", ("part", [c(LG, f"Rés princ {e}")], RPRINC)) for e in EPOQUES],
    ("Logement", "Part des ménages emménagés depuis moins de 2 ans (%)",
     ("part", [c(LG, "Ménages emménagés moins 2 ans")], c(LG, "Ménages"))),
    ("Logement", "Part des ménages emménagés depuis moins de 5 ans (%)",
     ("part", [c(LG, "Ménages emménagés moins 2 ans"), c(LG, "Ménages emménagés entre 2-4 ans")], c(LG, "Ménages"))),
    ("Logement", "Part des ménages emménagés depuis 10 ans ou plus (%)",
     ("part", [c(LG, "Ménages emménagés depuis 10 ans ou plus")], c(LG, "Ménages"))),
    ("Logement", "Part des ménages avec une voiture (%)", ("part", [c(LG, "Ménages une voiture")], c(LG, "Ménages"))),
    ("Logement", "Part des ménages avec au moins une voiture (%)", ("part", [c(LG, "Ménages au moins une voiture")], c(LG, "Ménages"))),
    ("Logement", "Part des RP suroccupées (%)",
     ("part", [c(LG, "résidences principales en suroccupation modérée"), c(LG, "résidences principales en suroccupation accentuée")], RPRINC)),
    # ------------------------------------------------------------------ Équipements (BPE)
    *[("Équipements", g, ("somme", c(BE, g))) for g in ["Écoles maternelles et élémentaires", "Collèges", "Lycées",
                                                       "Crèches et micro-crèches"]],
    ("Équipements", "Équipements culturels", ("somme", c(BE, "Équipements culturels"))),
    ("Équipements", "Équipements sportifs", ("somme", c(BE, "Équipements sportifs"))),
    ("Équipements", "dont gymnases et salles multisports", ("somme", c(BE, "Gymnases et salles multisports"))),
    ("Équipements", "dont plateaux et terrains extérieurs", ("somme", c(BE, "Plateaux et terrains extérieurs"))),
    *[("Équipements", g, ("somme", c(BE, g))) for g in ["Centres sociaux", "Services publics", "Bureaux et relais de poste",
                                                       "Commerces alimentaires de proximité"]],
    ("Équipements", "Supermarchés et hypermarchés", ("somme", c(BE, "Supermarchés et hypermarchés"))),
]


def colonnes_utiles():
    cols = set()
    for _, _, regle in INDICATEURS:
        for a in regle[1:]:
            if isinstance(a, list):
                cols.update(a)
            elif isinstance(a, str):
                cols.add(a)
    return cols


# --------------------------------------------------------------------------- lecture
# Nom de la source affiché dans les en-têtes, avec le millésime lu dans le nom d'onglet du TAB
SOURCES = {P: "RP", AR: "RP", CF: "RP", DF: "RP", LG: "RP", RD: "Filosofi", CA: "CAF", CN: "CNAM",
           DE: "France Travail", RP: "RPLS", BE: "BPE"}
MILLESIMES = {}
BPE_HORS_IRIS = pd.Series(dtype=float)     # équipements de Roubaix sans IRIS (« 59512_IND ») : comptés dans Roubaix seulement
NB_BPE_HORS_IRIS = 0                       # nombre de ces équipements appartenant à un groupe de BPE_GROUPES


def source(regle):
    """« RP 2023 », « Filosofi 2021 »... d'après la 1re colonne utilisée par l'indicateur."""
    col = regle[1][0] if isinstance(regle[1], list) else regle[1]
    alias = col.split("|", 1)[0]
    return f"{SOURCES[alias]} {MILLESIMES[alias]}" if alias in SOURCES else ""


def entete(lib, regle):
    src = source(regle)
    return f"{lib} [{src}]" if src else lib


def lire_donnees(codes):
    """DataFrame indexé par code IRIS, colonnes 'ONGLET|colonne'."""
    geo = json.loads(FICHIER_GEOJSON.read_text(encoding="utf-8"))
    nom_vers_code = {norm(f["properties"]["nom_iris"]): str(f["properties"]["code_iris"]) for f in geo["features"]}
    besoin = colonnes_utiles()
    xls = pd.ExcelFile(FICHIER_BDDMAS)
    morceaux = []
    for nom in xls.sheet_names:
        alias = " ".join(m for m in nom.split() if not m.isdigit())
        MILLESIMES[alias] = " ".join(m for m in nom.split() if m.isdigit())
        voulues = {b.split("|", 1)[1]: b for b in besoin if b.split("|", 1)[0] == alias}
        if not voulues:
            continue
        df = pd.read_excel(xls, nom, header=1)
        df = df[df["Libellé de l'Iris"].notna()]
        df.index = df["Libellé de l'Iris"].map(norm).map(nom_vers_code)
        colonnes = {norm(col): col for col in df.columns}
        for col, cle in voulues.items():
            if norm(col) not in colonnes:
                raise KeyError(f"Colonne '{col}' introuvable dans l'onglet '{nom}'")
            s = df[colonnes[norm(col)]].replace(VALEURS_SECRET, pd.NA)
            morceaux.append(pd.to_numeric(s, errors="coerce").rename(cle))
        log.info("Onglet %-20s : %d colonnes utilisées", nom, len(voulues))
    if any(b.startswith(BE + "|") for b in besoin):
        morceaux += [s for _, s in lire_bpe(codes).items()]
    if any(b.startswith("STAT|") for b in besoin):
        brut = pd.read_excel(FICHIER_STAT, sheet_name=0, header=None)
        stat = brut.iloc[2:].copy()
        stat.columns = brut.iloc[1].tolist()
        stat.index = stat["code_iris"].astype(str)
        for b in besoin:
            if b.startswith("STAT|"):
                morceaux.append(pd.to_numeric(stat[b.split("|", 1)[1]], errors="coerce").rename(b))
    donnees = pd.concat(morceaux, axis=1).reindex(codes)
    if donnees.index.isna().any() or len(donnees) != 38:
        raise ValueError("Les 38 IRIS ne sont pas tous retrouvés dans le TAB")
    return geo, donnees


def lire_bpe(codes):
    """Nombre d'équipements de chaque groupe de BPE_GROUPES par IRIS (colonnes 'BPE|groupe')."""
    global BPE_HORS_IRIS, NB_BPE_HORS_IRIS
    bpe = pd.read_excel(FICHIER_BPE, sheet_name="BPE", dtype=str)
    MILLESIMES[BE] = str(bpe["Année"].dropna().iloc[0])
    types = bpe["Code Type d'équipement"].fillna("")
    comptes = pd.DataFrame({c(BE, g): bpe[types.str.startswith(tuple(codes_bpe))].groupby("Code IRIS").size()
                            for g, codes_bpe in BPE_GROUPES.items()}).fillna(0)
    BPE_HORS_IRIS = comptes.drop(index=codes, errors="ignore").sum()
    retenus = types.str.startswith(tuple(code for liste in BPE_GROUPES.values() for code in liste))
    NB_BPE_HORS_IRIS = int((retenus & ~bpe["Code IRIS"].isin(codes)).sum())
    log.info("BPE %s : %d équipements, dont %d sans IRIS (comptés dans Roubaix seulement)", MILLESIMES[BE], len(bpe),
             (~bpe["Code IRIS"].isin(codes)).sum())
    return comptes.reindex(codes, fill_value=0)


def lire_cq(codes):
    cq = yaml.safe_load(FICHIER_CQ.read_text(encoding="utf-8"))
    groupe = pd.Series(HORS_CQ, index=codes)
    for nom, liste in cq.items():
        for code in map(str, liste):
            if code not in groupe.index:
                raise KeyError(f"{nom} : code IRIS inconnu {code}")
            if groupe[code] != HORS_CQ:
                raise ValueError(f"IRIS {code} rattaché à deux CQ : {groupe[code]} et {nom}")
            groupe[code] = nom
    return list(cq), groupe


# --------------------------------------------------------------------------- calcul
def calculer(donnees, groupe, ordre):
    zones = {nom: groupe.index[groupe == nom] for nom in ordre + [HORS_CQ]}
    zones[ROUBAIX] = groupe.index
    lignes = {}
    total = donnees.sum(min_count=1)
    for zone, iris in zones.items():
        d = donnees.loc[iris]
        s = d.sum(min_count=1)
        if zone == ROUBAIX and not BPE_HORS_IRIS.empty:
            s[BPE_HORS_IRIS.index] += BPE_HORS_IRIS
        res = {}
        for theme, lib, regle in INDICATEURS:
            t = regle[0]
            if t == "somme":
                v = s[regle[1]]
            elif t == "part":
                v = sum(s[a] for a in regle[1]) / s[regle[2]] * 100
            elif t == "part_diff":
                v = (s[regle[1]] - s[regle[2]]) / s[regle[3]] * 100
            elif t == "ratio":
                v = s[regle[1]] / s[regle[2]] * regle[3]
            elif t == "poids":
                v = s[regle[1]] / total[regle[1]] * 100
            elif t == "evol":
                v = (s[regle[1]] / s[regle[2]] - 1) * 100
            elif t == "approche":
                ok = d[regle[1]].notna()
                v = (d.loc[ok, regle[1]] * d.loc[ok, regle[2]]).sum() / d.loc[ok, regle[2]].sum() if ok.any() else float("nan")
            res[lib] = v
        res["Nombre d'IRIS"] = len(iris)
        lignes[zone] = res
    return zones, pd.DataFrame(lignes).T


def arrondir(lib, v):
    if v is None or (isinstance(v, float) and math.isnan(v)) or v is pd.NA:
        return None
    if lib.startswith(("Taille moyenne", "Superficie")):
        return round(float(v), 2)
    if "(%)" in lib or "(approché)" in lib or lib.startswith(("Densité", "Indice")):
        return round(float(v), 1) if "(€)" not in lib else int(round(float(v)))
    return int(round(float(v)))


# --------------------------------------------------------------------------- écriture
ENTETE = PatternFill("solid", fgColor="1F4E78")
TOTAL = PatternFill("solid", fgColor="D9E1F2")
METHODES = {"somme": "Somme des IRIS", "part": "Somme des effectifs / somme du dénominateur x 100",
            "part_diff": "(A - B) / dénominateur x 100, sur les sommes", "ratio": "Rapport des sommes",
            "poids": "Somme des IRIS de la zone / somme des 38 IRIS de Roubaix x 100",
            "evol": "Évolution des sommes", "approche": "Moyenne des IRIS pondérée par la population des ménages "
            "(revenus Filosofi non additifs ; IRIS secrétisés exclus)"}


def lire_commune():
    if not FICHIER_COMMUNE.exists():
        return {"source": "", "indicateurs": {}}
    return yaml.safe_load(FICHIER_COMMUNE.read_text(encoding="utf-8"))


# Pages récap (TOUT_ROUBAIX puis une page RECAP_<CQ> par Comité de Quartier) :
# (thème, libellé affiché, libellé de INDICATEURS ou de roubaix_commune.yml, format)
MODELE_ACCUEIL = "C"          # "A" tableau de bord, "B" sommaire latéral, "C" colonnes thématiques (scripts/accueil.py)
PAGE_ROUBAIX, PREFIXE_RECAP = "TOUT_ROUBAIX", "RECAP_"
POIDS = "Part de la population de Roubaix (%)"            # affiché sur les pages CQ seulement
CHIFFRES_CLES = [
    ("Démographie", "Habitants", "Population", "entier"),
    ("Démographie", "Poids dans la population de Roubaix", POIDS, "pct"),
    ("Démographie", "Évolution depuis 2016", "Évolution de la population depuis 2016 (%)", "pct"),
    ("Démographie", "Habitants au km²", "Densité (hab/km²)", "entier"),
    ("Démographie", "Moins de 20 ans", "Part des moins de 20 ans (%)", "pct"),
    ("Démographie", "65 ans ou plus", "Part des 65 ans ou plus (%)", "pct"),
    ("Démographie", "Femmes", "Part des femmes (%)", "pct"),
    ("Démographie", "Hommes", "Part des hommes (%)", "pct"),
    ("Démographie", "Ménages", "Ménages", "entier"),
    ("Démographie", "Taille moyenne des ménages", "Taille moyenne des ménages", "dec2"),
    ("Démographie", "Familles monoparentales", "Part des familles monoparentales (%)", "pct"),
    ("Socio-économie", "Niveau de vie médian", "Niveau de vie médian (€)", "euro"),
    ("Socio-économie", "Taux de pauvreté", "Taux de pauvreté au seuil de 60 % (%)", "pct"),
    ("Socio-économie", "Actifs (15-64 ans)", "Actifs 15-64 ans", "entier"),
    ("Socio-économie", "Inactifs (15-64 ans)", "Inactifs 15-64 ans", "entier"),
    ("Socio-économie", "Taux d'emploi (15-64 ans)", "Taux d'emploi des 15-64 ans (%)", "pct"),
    ("Socio-économie", "Taux de chômage (15-64 ans)", "Taux de chômage des 15-64 ans (%)", "pct"),
    ("Socio-économie", "Sans diplôme ou CEP (15 ans ou plus)", "Part des non-scolarisés de 15 ans ou plus sans diplôme ou CEP (%)", "pct"),
    ("Socio-économie", "Population couverte par la CAF", "Part de la population couverte par la CAF (%)", "pct"),
    ("Socio-économie", "Allocataires CAF au RSA socle", "Part des allocataires au RSA socle (%)", "pct"),
    ("Socio-économie", "Demandeurs d'emploi (cat. ABC)", "Demandeurs d'emploi cat. ABC", "entier"),
    ("Socio-économie", "Demandeurs d'emploi ABC / 15-64 ans", "Part des demandeurs d'emploi ABC parmi les 15-64 ans (%)", "pct"),
    ("Logement", "Logements", "Logements", "entier"),
    ("Logement", "Maisons", "Maisons", "entier"),
    ("Logement", "Appartements", "Appartements", "entier"),
    ("Logement", "Logements vacants", "Part des logements vacants (%)", "pct"),
    ("Logement", "Locataires HLM", "Part des locataires HLM (%)", "pct"),
    ("Logement", "Logements sociaux", "Logements sociaux (RPLS)", "entier"),
    ("Logement", "Résidences construites avant 1946", None, "pct"),
    ("Logement", "Ménages avec une voiture", "Part des ménages avec une voiture (%)", "pct"),
    ("Équipements", "Écoles maternelles et élémentaires", "Écoles maternelles et élémentaires", "entier"),
    ("Équipements", "Collèges", "Collèges", "entier"),
    ("Équipements", "Lycées", "Lycées", "entier"),
    ("Équipements", "Crèches et micro-crèches", "Crèches et micro-crèches", "entier"),
    ("Équipements", "Équipements culturels", "Équipements culturels", "entier"),
    ("Équipements", "Équipements sportifs", "Équipements sportifs", "entier"),
    ("Équipements", "Centres sociaux", "Centres sociaux", "entier"),
    ("Équipements", "Services publics", "Services publics", "entier"),
    ("Équipements", "Bureaux et relais de poste", "Bureaux et relais de poste", "entier"),
    ("Équipements", "Commerces alimentaires de proximité", "Commerces alimentaires de proximité", "entier"),
    ("Équipements", "Supermarchés", "Supermarchés et hypermarchés", "entier"),
]
# Indicateur communal (roubaix_commune.yml) -> équivalent par IRIS, utilisé pour les CQ
EQUIVALENTS_IRIS = {
    "Niveau de vie médian (€)": "Revenu disponible médian par UC (€) (approché)",
    "Taux de pauvreté au seuil de 60 % (%)": "Taux de pauvreté au seuil de 60 % (%) (approché)",
}
PAGES = [
    ("Démographie", "Population, densité, évolution depuis 2006 et 2016, tranches d'âge, genre, "
                    "taille et composition des ménages, familles monoparentales et nombreuses."),
    ("Socio-économie", "Revenus et pauvreté, activité, emploi, chômage, catégories socioprofessionnelles, "
                       "diplômes, précarité (CAF, RSA, AAH, C2S, demandeurs d'emploi)."),
    ("Logement", "Parc, vacance, statut d'occupation, logements sociaux, typologie, "
                 "époque de construction, ancienneté d'emménagement, suroccupation."),
    ("Équipements", "Écoles, crèches, culture, sport, centres sociaux, services publics et commerces "
                    "de proximité (nombre d'équipements)."),
    ("Méthode", "Source, millésime et mode de calcul de chaque indicateur."),
    ("Rattachement IRIS", "Les 38 IRIS de Roubaix et leur Comité de Quartier."),
]
# Quartiers vécus contenus dans un IRIS (non découpables : l'IRIS est le plus petit niveau Insee).
# Affichés « IRIS (contient ...) » dans la liste des CQ de la page d'accueil.
QUARTIERS_CONTENUS = {
    "595120603": "Motte, Bossut",      # Anseele
    "595120601": "Crouy",              # Espérance Centre
    "595121103": "Carihem",            # Sartel
    "595121201": "Linné Sud",          # Linné-Chemin Neuf
    "595121302": "Petites Haies",      # Hauts Champs
}
DESCRIPTIONS_SOURCES = {
    "RP": "Recensement de la population (Insee), par IRIS",
    "Filosofi": "Revenus localisés sociaux et fiscaux (Insee), dernier millésime publié par IRIS",
    "CAF": "Allocataires de la Caisse d'allocations familiales, par IRIS",
    "CNAM": "Assurés et bénéficiaires de la C2S (Assurance maladie), par IRIS",
    "France Travail": "Demandeurs d'emploi en fin de mois, par IRIS",
    "RPLS": "Logements des bailleurs sociaux, par IRIS",
    "BPE": "Base permanente des équipements (Insee), équipements géolocalisés à l'IRIS",
}


def nom_recap(cq):
    """« CQ Sainte Elisabeth » -> « RECAP_Sainte_Elisabeth » (nom d'onglet Excel : 31 caractères, sans : \\ / ? * [ ])."""
    nom = cq.removeprefix("CQ ").strip().replace(" ", "_")
    nom = "".join(ch for ch in nom if ch not in ':\\/?*[]')
    return (PREFIXE_RECAP + nom)[:31]


def chiffres_cles(zone, table, commune):
    regles = {lib: r for _, lib, r in INDICATEURS}
    ligne = table.loc[zone]
    kpis = []
    for theme, court, lib, fmt in CHIFFRES_CLES:
        if lib is None:          # RP construites avant 1946 = avant 1919 + 1919-1945
            libs = ["Part des RP construites avant 1919 (%)", "Part des RP construites 1919 à 1945 (%)"]
            kpis.append((theme, court, arrondir(libs[0], sum(float(ligne[x]) for x in libs)), fmt, source(regles[libs[0]])))
        elif lib == POIDS and zone == ROUBAIX:
            continue
        elif zone == ROUBAIX and lib in commune["indicateurs"]:
            kpis.append((theme, court, commune["indicateurs"][lib], fmt, commune["source"]))
        else:
            lib_iris = EQUIVALENTS_IRIS.get(lib, lib)
            court_iris = court + " (approché)" if lib_iris != lib else court
            kpis.append((theme, court_iris, arrondir(lib_iris, ligne[lib_iris]), fmt, source(regles[lib_iris]) or "RP"))
    return kpis


def ecrire_accueil(wb, table, zones, ordre, noms_iris, commune, modele=None):
    from accueil import MODELES
    dessiner = MODELES[modele or MODELE_ACCUEIL]
    sources = [(s, DESCRIPTIONS_SOURCES[s.rsplit(" ", 1)[0]]) for s in sorted({source(r) for _, _, r in INDICATEURS if source(r)})]
    if commune["indicateurs"]:
        sources.append((commune["source"], "Niveau de vie médian et taux de pauvreté de la commune (Insee), non publiés par IRIS"))

    def iris_de(cq):
        return ", ".join(noms_iris[c] + (f" (contient {QUARTIERS_CONTENUS[c]})" if c in QUARTIERS_CONTENUS else "")
                         for c in zones[cq])

    note_defm = "Demandeurs d'emploi / 15-64 ans : DEFM France Travail rapportés à la population du recensement (millésimes différents)."
    note_bpe = ("Équipements : nombre d'établissements recensés par la BPE, sans leur taille (places, surfaces). "
                "Les espaces verts ne figurent pas dans la BPE.")
    dessiner(wb, {
        "titre": f"Portrait des Comités de Quartier de Roubaix · {ANNEE}",
        "sous_titre": f"Chiffres clés de Roubaix. Une page récap par Comité de Quartier (CQ) : liens à droite. "
                      "Chaque indicateur indique sa source et son année.",
        "pages": PAGES,
        "kpis": chiffres_cles(ROUBAIX, table, commune),
        "sources": sources,
        "notes": [
            "Parts et taux des CQ recalculés sur les effectifs additionnés de leurs IRIS.",
            "Revenus par CQ « approchés » (Filosofi 2021, moyenne pondérée) ; vides pour BEH Alma (secret statistique).",
            "Filosofi 2023 (commune) n'est pas comparable à Filosofi 2021 (IRIS) : changement de méthode Insee.",
            note_defm,
            note_bpe + f" {NB_BPE_HORS_IRIS} équipements sans IRIS connu sont comptés dans Roubaix mais dans aucun CQ.",
            f"Roubaix = 38 IRIS, dont {len(zones[HORS_CQ])} hors CQ ({', '.join(noms_iris[c] for c in zones[HORS_CQ])}).",
        ],
        "cqs": [(cq, iris_de(cq), nom_recap(cq)) for cq in ordre],
    }, nom=PAGE_ROUBAIX, index=0)
    sources_cq = [s for s in sources if s[0] != commune["source"]]
    for cq in ordre:                       # pages récap en fin de classeur, après Rattachement IRIS
        dessiner(wb, {
            "titre": f"{cq} · {ANNEE}",
            "sous_titre": f"IRIS : {iris_de(cq)}.",
            "pages": PAGES,
            "kpis": chiffres_cles(cq, table, commune),
            "sources": sources_cq,
            "notes": [
                "Parts et taux recalculés sur les effectifs additionnés des IRIS du CQ.",
                "Revenus « approchés » : Filosofi 2021, moyenne des IRIS pondérée par la population des ménages.",
                "Poids : population du CQ / population de Roubaix (38 IRIS).",
                note_defm,
                note_bpe,
            ],
            "cqs": [],
            "retour": PAGE_ROUBAIX,
        }, nom=nom_recap(cq), index=None)


def ecrire_excel(table, zones, groupe, noms_iris, ordre):
    chemin = TRAITE / f"portrait_cq_{ANNEE}.xlsx"
    commune = lire_commune()
    wb = Workbook()
    wb.remove(wb.active)
    for theme in ("Démographie", "Socio-économie", "Logement", "Équipements"):
        libs = [lib for th, lib, _ in INDICATEURS if th == theme]
        entetes = [entete(lib, r) for th, lib, r in INDICATEURS if th == theme]
        ws = wb.create_sheet(theme)
        extra = commune["indicateurs"] if theme == "Socio-économie" else {}
        entetes += [f"{lib} [{commune['source']}]" for lib in extra]
        ws.append(["Comité de quartier", "Nombre d'IRIS"] + entetes)
        for zone, row in table.iterrows():
            if zone == HORS_CQ and not AFFICHER_HORS_CQ:
                continue
            valeurs = [arrondir(lib, row[lib]) for lib in libs]
            valeurs += [v if zone == ROUBAIX else None for v in extra.values()]
            ws.append([zone, int(row["Nombre d'IRIS"])] + valeurs)
        for cell in ws[1]:
            cell.font, cell.fill = Font(bold=True, color="FFFFFF"), ENTETE
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        for cell in ws[ws.max_row]:
            cell.font, cell.fill = Font(bold=True), TOTAL
        ws.row_dimensions[1].height = 75
        ws.column_dimensions["A"].width = 22
        for j in range(2, len(libs) + 3):
            ws.column_dimensions[get_column_letter(j)].width = 15
        ws.freeze_panes = "C2"
    ws = wb.create_sheet("Méthode")
    ws.append(["Thème", "Indicateur", "Source", "Calcul", "Numérateur / colonne", "Dénominateur"])
    for theme, lib, regle in INDICATEURS:
        t = regle[0]
        num = " + ".join(regle[1]) if isinstance(regle[1], list) else regle[1]
        if t == "part_diff":
            num = f"{regle[1]} - {regle[2]}"
        den = "" if t == "somme" else f"{regle[1]} (Roubaix)" if t == "poids" else regle[3] if t == "part_diff" else regle[2]
        if t == "approche":
            den = f"pondération : {den}"
        if num.startswith(BE + "|"):
            num += " (types " + ", ".join(BPE_GROUPES[num.split("|", 1)[1]]) + ")"
        ws.append([theme, lib, source(regle) or "RP (stat_insee_iris)", METHODES[t], num, den])
    ws = wb.create_sheet("Rattachement IRIS")
    ws.append(["Code IRIS", "IRIS", "Comité de quartier"])
    for code in sorted(groupe.index):
        ws.append([code, noms_iris[code], groupe[code]])
    for w in wb.worksheets[-2:]:
        for cell in w[1]:
            cell.font, cell.fill = Font(bold=True, color="FFFFFF"), ENTETE
        for j, larg in enumerate([16, 60, 18, 50, 70, 40], start=1):
            w.column_dimensions[get_column_letter(j)].width = larg
        w.auto_filter.ref = w.dimensions
    ecrire_accueil(wb, table, zones, ordre, noms_iris, commune)
    wb.save(chemin)
    log.info("Excel écrit : %s", chemin)


def ecrire_geojson(geo, table, zones):
    chemin = TRAITE / f"cq_roubaix_{ANNEE}.geojson"
    formes = {str(f["properties"]["code_iris"]): shape(f["geometry"]) for f in geo["features"]}
    feats = []
    for zone, row in table.iterrows():
        if zone == ROUBAIX or (zone == HORS_CQ and not AFFICHER_HORS_CQ):
            continue
        poly = unary_union([formes[code].buffer(0) for code in zones[zone]])
        props = {"cq": zone, "annee": ANNEE, "nb_iris": int(row["Nombre d'IRIS"]),
                 "iris": ", ".join(sorted(zones[zone]))}
        props.update({entete(lib, r): arrondir(lib, row[lib]) for _, lib, r in INDICATEURS})
        feats.append({"type": "Feature", "properties": props, "geometry": mapping(poly)})
    sortie = {"type": "FeatureCollection", "name": f"cq_roubaix_{ANNEE}", "crs": geo.get("crs"), "features": feats}
    chemin.write_text(json.dumps(sortie, ensure_ascii=False), encoding="utf-8")
    log.info("GeoJSON écrit : %s (%d entités)", chemin, len(feats))


def main():
    global FICHIER_BDDMAS, FICHIER_GEOJSON, FICHIER_STAT, FICHIER_BPE
    FICHIER_BDDMAS = chercher(f"TAB_BDDMAS_Tout_perimetre_{ANNEE}.xlsx", "brut")
    FICHIER_GEOJSON = chercher("Iris_2000.geojson", "brut")
    FICHIER_STAT = chercher(f"stat_insee_iris_{ANNEE}.xlsx", "traite", "brut")   # populations 2006 / 2011 / 2016
    FICHIER_BPE = chercher_bpe()
    log.info("Sources : %s | %s | %s | %s", FICHIER_BDDMAS, FICHIER_GEOJSON, FICHIER_STAT, FICHIER_BPE)
    TRAITE.mkdir(parents=True, exist_ok=True)
    geo = json.loads(FICHIER_GEOJSON.read_text(encoding="utf-8"))
    codes = sorted(str(f["properties"]["code_iris"]) for f in geo["features"])
    noms_iris = {str(f["properties"]["code_iris"]): f["properties"]["nom_iris"] for f in geo["features"]}
    ordre, groupe = lire_cq(codes)
    geo, donnees = lire_donnees(codes)
    zones, table = calculer(donnees, groupe, ordre)
    log.info("Hors CQ : %s", ", ".join(noms_iris[c] for c in zones[HORS_CQ]) or "aucun")
    ecrire_excel(table, zones, groupe, noms_iris, ordre)
    ecrire_geojson(geo, table, zones)


if __name__ == "__main__":
    main()
