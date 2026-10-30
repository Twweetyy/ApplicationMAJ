"""
Modèles de page d'accueil pour portrait_cq.py.

Chaque modèle reçoit le classeur et un dictionnaire de contenu :
  titre, sous_titre                      : textes
  pages   : [(nom_feuille, description)]
  kpis    : [(theme, libellé court, valeur, format, source)]   theme = Démographie / Socio-économie / Logement / Équipements
  sources : [(source, description)]
  notes   : [texte]
  cqs     : [(cq, "IRIS, IRIS...", onglet_recap)]   lien vers la page récap du CQ
  retour  : onglet vers lequel pointe le lien « Retour » (pages récap des CQ), facultatif
Chaque modèle crée l'onglet `nom` à la position `index` (None : en dernier).
Le modèle utilisé se choisit avec MODELE_ACCUEIL dans portrait_cq.py ("A", "B" ou "C").
"""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BLEU, BLEU_CLAIR, GRIS, GRIS_CLAIR, BLANC = "1F4E78", "DDEBF7", "595959", "F2F2F2", "FFFFFF"
COULEURS = {"Démographie": "2E75B6", "Socio-économie": "C55A11", "Logement": "548235", "Équipements": "7030A0"}
CLAIRS = {"Démographie": "DEEAF6", "Socio-économie": "FBE5D6", "Logement": "E2F0D9", "Équipements": "EADCF4"}
FORMATS = {"entier": '#,##0', "pct": '0.0" %"', "euro": '#,##0" €"', "dec1": '#,##0.0', "dec2": '0.00'}
FIN = Side(style="thin", color="BFBFBF")


def _fill(couleur):
    return PatternFill("solid", fgColor=couleur)


def _bloc(ws, plage, valeur=None, font=None, fill=None, align=None, lien=None):
    """Fusionne une plage, écrit la valeur et applique le style à toutes ses cellules."""
    ws.merge_cells(plage)
    cellule = ws[plage.split(":")[0]]
    cellule.value = valeur
    for ligne in ws[plage]:
        for c in ligne:
            if fill:
                c.fill = fill
    if lien:                       # le lien d'abord : il impose le style « Lien hypertexte », la police vient après
        cellule.hyperlink = f"#'{lien}'!A1"
    if font:
        cellule.font = font
    cellule.alignment = align or Alignment(vertical="center", wrap_text=True)
    return cellule


def _cadre(ws, plage, couleur_gauche=None):
    lignes = list(ws[plage])
    for i, ligne in enumerate(lignes):
        for j, c in enumerate(ligne):
            c.border = Border(top=FIN if i == 0 else None, bottom=FIN if i == len(lignes) - 1 else None,
                              left=Side(style="thick", color=couleur_gauche) if (j == 0 and couleur_gauche) else (FIN if j == 0 else None),
                              right=FIN if j == len(ligne) - 1 else None)


def _lien(cellule, onglet, couleur=BLEU, gras=True, taille=None):
    cellule.hyperlink = f"#'{onglet}'!A1"
    cellule.font = Font(bold=gras, color=couleur, underline="single", size=taille)


def _retour(ws, d, cellule):
    if d.get("retour"):
        ws[cellule].value = f"← Retour à {d['retour']}"
        _lien(ws[cellule], d["retour"])


def _grille(ws, largeurs):
    ws.sheet_view.showGridLines = False
    for i, l in enumerate(largeurs, start=1):
        ws.column_dimensions[get_column_letter(i)].width = l


# --------------------------------------------------------------------------- A : tableau de bord
def modele_a(wb, d, nom="Accueil", index=0):
    """Bandeau titre, boutons de navigation, tuiles de chiffres clés colorées par thème."""
    ws = wb.create_sheet(nom, index)
    _grille(ws, [2] + [11] * 12 + [2])
    _bloc(ws, "B2:M3", d["titre"], Font(bold=True, size=20, color=BLANC), _fill(BLEU),
          Alignment(vertical="center", indent=1))
    _bloc(ws, "B4:M4", d["sous_titre"], Font(italic=True, size=10, color=GRIS))
    ws.row_dimensions[4].height = 30
    _retour(ws, d, "B5")

    r = 6
    _bloc(ws, f"B{r}:M{r}", "NAVIGUER", Font(bold=True, size=9, color=GRIS))
    r += 1
    for i, (page, _) in enumerate(d["pages"]):
        col = 2 + i * 12 // len(d["pages"])
        fin = 1 + (i + 1) * 12 // len(d["pages"])
        couleur = COULEURS.get(page, GRIS)
        plage = f"{get_column_letter(col)}{r}:{get_column_letter(fin)}{r + 1}"
        _bloc(ws, plage, page, Font(bold=True, color=BLANC, size=11), _fill(couleur),
              Alignment(horizontal="center", vertical="center"), lien=page)
    r += 3

    _bloc(ws, f"B{r}:M{r}", "ROUBAIX EN CHIFFRES", Font(bold=True, size=9, color=GRIS))
    r += 1
    for k, (theme, lib, val, fmt, src) in enumerate(d["kpis"]):
        ligne, pos = divmod(k, 4)
        haut = r + ligne * 5
        c1, c2 = get_column_letter(2 + pos * 3), get_column_letter(3 + pos * 3)
        fond = _fill(CLAIRS.get(theme, GRIS_CLAIR))
        v = _bloc(ws, f"{c1}{haut}:{c2}{haut + 1}", val, Font(bold=True, size=20, color=COULEURS.get(theme, BLEU)), fond,
                  Alignment(horizontal="left", vertical="center", indent=1))
        v.number_format = FORMATS[fmt]
        _bloc(ws, f"{c1}{haut + 2}:{c2}{haut + 2}", lib, Font(size=10, bold=True, color="262626"), fond,
              Alignment(vertical="top", wrap_text=True, indent=1))
        _bloc(ws, f"{c1}{haut + 3}:{c2}{haut + 3}", src, Font(size=8, color=GRIS), fond, Alignment(indent=1))
        _cadre(ws, f"{c1}{haut}:{c2}{haut + 3}", COULEURS.get(theme))
        ws.row_dimensions[haut + 2].height = 28
    r += (len(d["kpis"]) + 3) // 4 * 5 + 1

    _bloc(ws, f"B{r}:G{r}", "À SAVOIR", Font(bold=True, size=9, color=GRIS))
    _bloc(ws, f"H{r}:M{r}", "SOURCES", Font(bold=True, size=9, color=GRIS))
    r += 1
    for i in range(max(len(d["notes"]), len(d["sources"]))):
        if i < len(d["notes"]):
            _bloc(ws, f"B{r + i}:G{r + i}", "• " + d["notes"][i], Font(size=9), _fill(GRIS_CLAIR))
        if i < len(d["sources"]):
            s, t = d["sources"][i]
            _bloc(ws, f"H{r + i}:I{r + i}", s, Font(size=9, bold=True), _fill(GRIS_CLAIR))
            _bloc(ws, f"J{r + i}:M{r + i}", t, Font(size=9, color=GRIS), _fill(GRIS_CLAIR))
        ws.row_dimensions[r + i].height = 42
    r += max(len(d["notes"]), len(d["sources"])) + 1

    if d["cqs"]:
        _bloc(ws, f"B{r}:M{r}", "LES COMITÉS DE QUARTIER (cliquer pour la page récap)", Font(bold=True, size=9, color=GRIS))
    r += 1
    for i, (cq, iris, recap) in enumerate(d["cqs"]):
        ligne, pos = divmod(i, 2)
        c = 2 + pos * 6
        _bloc(ws, f"{get_column_letter(c)}{r + ligne}:{get_column_letter(c + 1)}{r + ligne}", cq, Font(bold=True, size=9))
        _lien(ws.cell(r + ligne, c), recap, taille=9)
        _bloc(ws, f"{get_column_letter(c + 2)}{r + ligne}:{get_column_letter(c + 5)}{r + ligne}", iris, Font(size=9, color=GRIS))
        ws.row_dimensions[r + ligne].height = 26


# --------------------------------------------------------------------------- B : sommaire latéral
def modele_b(wb, d, nom="Accueil", index=0):
    """Colonne de sommaire à gauche, contenu en tableau sobre à droite."""
    ws = wb.create_sheet(nom, index)
    _grille(ws, [2, 26, 3, 52, 16, 26, 2])
    hauteur = 12 + len(d["kpis"]) + len(d["notes"]) + len(d["sources"]) + len(d["cqs"]) + 12
    for r in range(1, hauteur):
        ws.cell(r, 2).fill = _fill(BLEU)
    ws["B2"].value, ws["B2"].font = "SOMMAIRE", Font(bold=True, color=BLANC, size=12)
    for i, (page, _) in enumerate(d["pages"]):
        c = ws.cell(4 + i * 2, 2, f"›  {page}")
        c.hyperlink = f"#'{page}'!A1"
        c.font = Font(bold=True, color=BLANC, size=11)

    ws["D2"].value, ws["D2"].font = d["titre"], Font(bold=True, size=18, color=BLEU)
    _retour(ws, d, "D4")
    _bloc(ws, "D3:F3", d["sous_titre"], Font(italic=True, size=10, color=GRIS))
    ws.row_dimensions[3].height = 30

    def titre(r, texte):
        _bloc(ws, f"D{r}:F{r}", texte, Font(bold=True, size=12, color=BLEU))
        for c in ws[f"D{r}:F{r}"][0]:
            c.border = Border(bottom=Side(style="medium", color=BLEU))

    r = 5
    titre(r, "Roubaix en chiffres")
    r += 1
    for j, t in enumerate(["Indicateur", "Valeur", "Source"]):
        c = ws.cell(r, 4 + j, t)
        c.font, c.fill = Font(bold=True, color=BLANC), _fill(GRIS)
        c.alignment = Alignment(horizontal="right" if j == 1 else "left")
    theme_prec = None
    for i, (theme, lib, val, fmt, src) in enumerate(d["kpis"]):
        r += 1
        if theme != theme_prec:
            c = ws.cell(r, 4, theme)
            c.font = Font(bold=True, color=COULEURS[theme], size=10)
            r += 1
            theme_prec = theme
        fond = _fill(GRIS_CLAIR) if i % 2 else None
        for j, v in enumerate([lib, val, src]):
            c = ws.cell(r, 4 + j, v)
            if fond:
                c.fill = fond
            c.font = Font(bold=(j == 1), size=10, color=GRIS if j == 2 else "000000")
            c.border = Border(bottom=Side(style="hair", color="D9D9D9"))
        ws.cell(r, 5).number_format = FORMATS[fmt]
    r += 2
    titre(r, "À savoir")
    for n in d["notes"]:
        r += 1
        _bloc(ws, f"D{r}:F{r}", "• " + n, Font(size=9))
        ws.row_dimensions[r].height = 30
    r += 2
    titre(r, "Sources")
    for s, t in d["sources"]:
        r += 1
        ws.cell(r, 4, s).font = Font(bold=True, size=9)
        _bloc(ws, f"E{r}:F{r}", t, Font(size=9, color=GRIS))
        ws.row_dimensions[r].height = 26
    r += 2
    if d["cqs"]:
        titre(r, "Les Comités de Quartier (cliquer pour la page récap)")
    for cq, iris, recap in d["cqs"]:
        r += 1
        _lien(ws.cell(r, 4, cq), recap, taille=9)
        _bloc(ws, f"E{r}:F{r}", iris, Font(size=9, color=GRIS))


# --------------------------------------------------------------------------- C : une colonne par thème
def modele_c(wb, d, nom="Accueil", index=0):
    """Une colonne par thème avec ses chiffres clés et un lien vers la page."""
    ws = wb.create_sheet(nom, index)
    _grille(ws, [2, 30, 12, 3, 30, 12, 3, 30, 12, 3, 30, 12, 2])
    _bloc(ws, "B2:L2", d["titre"], Font(bold=True, size=20, color=BLEU), align=Alignment(vertical="center"))
    ws.row_dimensions[2].height = 34
    _bloc(ws, "B3:L3", d["sous_titre"], Font(italic=True, size=10, color=GRIS))
    ws.row_dimensions[3].height = 30
    _retour(ws, d, "B4")
    descriptions = dict(d["pages"])
    hauteur_max = 0
    for k, theme in enumerate(COULEURS):
        c1, c2 = get_column_letter(2 + 3 * k), get_column_letter(3 + 3 * k)
        _bloc(ws, f"{c1}5:{c2}6", theme.upper(), Font(bold=True, size=14, color=BLANC), _fill(COULEURS[theme]),
              Alignment(horizontal="center", vertical="center"), lien=theme)
        _bloc(ws, f"{c1}7:{c2}7", descriptions.get(theme, ""), Font(size=9, color=GRIS), _fill(CLAIRS[theme]),
              Alignment(wrap_text=True, vertical="top"))
        ws.row_dimensions[7].height = 58
        r = 8
        for th, lib, val, fmt, src in d["kpis"]:
            if th != theme:
                continue
            a = ws[f"{c1}{r}"]
            a.value, a.font = lib, Font(size=10)
            a.alignment = Alignment(wrap_text=True, vertical="center")
            b = ws[f"{c2}{r}"]
            b.value, b.number_format = val, FORMATS[fmt]
            b.font = Font(bold=True, size=13, color=COULEURS[theme])
            b.alignment = Alignment(horizontal="right", vertical="center")
            s = ws[f"{c1}{r + 1}"]
            s.value, s.font = src, Font(size=8, color=GRIS)
            for c in (a, b):
                c.border = Border(top=Side(style="hair", color="BFBFBF"))
            ws.row_dimensions[r].height = 30
            r += 2
        lien = ws[f"{c1}{r}"]
        lien.value = f"→ Voir la page {theme}"
        lien.hyperlink = f"#'{theme}'!A1"
        lien.font = Font(bold=True, color=COULEURS[theme], underline="single")
        hauteur_max = max(hauteur_max, r)
    if d["cqs"]:                   # sommaire des pages récap : bloc séparé, à droite des quatre thèmes (colonnes N à P)
        ws.column_dimensions["M"].width = 8
        for col in "NOP":
            ws.column_dimensions[col].width = 22
        _bloc(ws, "N5:P6", "PAGES RÉCAP PAR COMITÉ DE QUARTIER", Font(bold=True, size=14, color=BLANC), _fill(BLEU),
              Alignment(horizontal="center", vertical="center"))
        _bloc(ws, "N7:P7", "Cliquer sur un CQ pour ouvrir sa page de chiffres clés.", Font(size=9, color=GRIS),
              _fill(BLEU_CLAIR), Alignment(wrap_text=True, vertical="top", indent=1))
        for i, (cq, _, recap) in enumerate(d["cqs"]):
            ligne, pos = divmod(i, 3)
            col = "NOP"[pos]
            c = _bloc(ws, f"{col}{8 + 2 * ligne}:{col}{9 + 2 * ligne}", f"→ {cq}", fill=_fill(BLEU_CLAIR),
                      align=Alignment(horizontal="center", vertical="center"))
            _lien(c, recap)
        fin = 8 + 2 * ((len(d["cqs"]) + 2) // 3) - 1
        for ligne in ws[f"N8:P{fin}"]:          # cases vides de la dernière ligne
            for c in ligne:
                c.fill = _fill(BLEU_CLAIR)
        _cadre(ws, f"N5:P{fin}")
    r = hauteur_max + 2
    autres = [p for p in d["pages"] if p[0] not in COULEURS]
    _bloc(ws, f"B{r}:L{r}", "Autres pages : " + "  ·  ".join(p for p, _ in autres), Font(size=10, color=BLEU))
    r += 2
    _bloc(ws, f"B{r}:C{r}", "À savoir", Font(bold=True, color=BLANC), _fill(GRIS))
    _bloc(ws, f"E{r}:F{r}", "Sources", Font(bold=True, color=BLANC), _fill(GRIS))
    if d["cqs"]:
        _bloc(ws, f"H{r}:I{r}", "Comités de Quartier", Font(bold=True, color=BLANC), _fill(GRIS))
    for i, n in enumerate(d["notes"]):
        _bloc(ws, f"B{r + 1 + i}:C{r + 1 + i}", "• " + n, Font(size=9))
    for i, (s, t) in enumerate(d["sources"]):
        _bloc(ws, f"E{r + 1 + i}:F{r + 1 + i}", f"{s} : {t}", Font(size=9))
    for i, (cq, iris, recap) in enumerate(d["cqs"]):
        _bloc(ws, f"H{r + 1 + i}:I{r + 1 + i}", f"{cq} : {iris}", Font(size=9))
        _lien(ws[f"H{r + 1 + i}"], recap, gras=False, taille=9)
    for i in range(max(len(d["notes"]), len(d["sources"]), len(d["cqs"]))):
        ws.row_dimensions[r + 1 + i].height = 40


MODELES = {"A": modele_a, "B": modele_b, "C": modele_c}
